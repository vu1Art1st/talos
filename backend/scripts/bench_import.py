"""导入性能基准（ROADMAP P1-5）：20 份报告上限下的耗时 / 内存 / 磁盘峰值记录。

用法（backend 目录，使用 .venv 解释器）::

    .venv/Scripts/python -m scripts.bench_import                    # 默认 20 份 × 12 漏洞
    .venv/Scripts/python -m scripts.bench_import --files 5 --vulns 6
    .venv/Scripts/python -m scripts.bench_import --out bench_import.json

口径：
- 生成的文档与「平台导出的渗透测试报告」同构（测试目标表 + 时间与人员表 + 风险问题汇总表 +
  风险问题详情章节），因此走的是与线上一致的 `parse_any_docx` 报告分支；
- 内存用 `tracemalloc` 峰值（仅统计 Python 侧分配，不含 docx 库的 C 侧开销，属保守下界）；
- 磁盘 = 生成的 docx 总字节 + 解析期落盘的图片字节；
- 单一进程顺序解析（与 worker 的导入任务一致），不并发。

输出：stdout 表格 + JSON（默认写入 `storage/bench_import.json`，便于随发布记录归档）。
"""
from __future__ import annotations

import argparse
import json
import shutil
import tempfile
import time
import tracemalloc
from pathlib import Path

_LEVELS = ("严重", "高危", "中危", "低危")
_TYPES = ("SQL注入漏洞", "信息泄露", "逻辑漏洞", "越权访问", "跨站脚本攻击")


def build_report_docx(system_name: str, vuln_count: int):
    """生成与平台导出一致的报告 docx（返回 Document 对象）。"""
    from docx import Document

    doc = Document()
    target = doc.add_table(rows=5, cols=2)
    target.rows[0].cells[0].text, target.rows[0].cells[1].text = "业务系统名称", system_name
    target.rows[1].cells[0].text, target.rows[1].cells[1].text = "被测系统URL", "https://bench.example.com"
    target.rows[2].cells[0].text, target.rows[2].cells[1].text = "被测系统域名", "bench.example.com"
    target.rows[3].cells[0].text, target.rows[3].cells[1].text = "被测系统IP", "192.0.2.20"
    target.rows[4].cells[0].text, target.rows[4].cells[1].text = "被测测试账号", "demo/demo"

    schedule = doc.add_table(rows=5, cols=4)
    schedule.rows[0].cells[0].text = "测试工作时间段"
    schedule.rows[1].cells[0].text, schedule.rows[1].cells[1].text = "起始时间", "2026-08-01"
    schedule.rows[1].cells[2].text, schedule.rows[1].cells[3].text = "结束时间", "2026-08-10"
    for i, head in enumerate(("参测人员", "所属部门", "人员角色", "人员分工")):
        schedule.rows[3].cells[i].text = head
    schedule.rows[4].cells[0].text = "张三"

    doc.add_heading("风险问题汇总", level=1)
    summary = doc.add_table(rows=1 + vuln_count, cols=4)
    for i, head in enumerate(("问题等级", "风险类型", "风险问题", "修复状态")):
        summary.rows[0].cells[i].text = head
    titles = [f"基准用例漏洞{i:03d}" for i in range(1, vuln_count + 1)]
    for index, title in enumerate(titles, start=1):
        summary.rows[index].cells[0].text = _LEVELS[index % len(_LEVELS)]
        summary.rows[index].cells[1].text = _TYPES[index % len(_TYPES)]
        summary.rows[index].cells[2].text = title
        summary.rows[index].cells[3].text = "已修复" if index % 3 == 0 else "未修复"

    doc.add_heading("风险问题详情", level=1)
    for index, title in enumerate(titles, start=1):
        doc.add_heading(f"{title}（未修复）", level=3)
        doc.add_paragraph("漏洞等级：" + _LEVELS[index % len(_LEVELS)])
        doc.add_paragraph("漏洞类型：" + _TYPES[index % len(_TYPES)])
        doc.add_paragraph(f"漏洞链接：https://bench.example.com/api/{index}")
        doc.add_paragraph("漏洞描述")
        doc.add_paragraph("基准生成的漏洞描述内容。" * 20)
        doc.add_paragraph("漏洞证明")
        doc.add_paragraph("基准生成的复现步骤内容。" * 20)
        doc.add_paragraph("修复建议")
        doc.add_paragraph("基准生成的修复建议内容。" * 20)
    return doc


def _dir_size(path: Path) -> int:
    return sum(f.stat().st_size for f in path.rglob("*") if f.is_file())


def run(files: int, vulns: int, out: Path) -> dict:
    from app.services.docx_parser import parse_any_docx

    tmp = Path(tempfile.mkdtemp(prefix="talos_bench_"))
    docs_dir = tmp / "docs"
    docs_dir.mkdir()
    images_dir = tmp / "images"
    images_dir.mkdir()

    try:
        paths: list[Path] = []
        for index in range(files):
            name = f"bench_report_{index:02d}.docx"
            path = docs_dir / name
            build_report_docx(f"基准系统{index:02d}", vulns).save(str(path))
            paths.append(path)
        gen_bytes = _dir_size(docs_dir)

        records_total = 0
        mismatches = 0
        tracemalloc.start()
        started = time.perf_counter()
        for path in paths:
            _kind, _meta, records = parse_any_docx(
                str(path), str(images_dir), "/storage/uploads/images", path.name,
            )
            records_total += len(records)
            mismatches += sum(1 for r in records if r.get("level_mismatch"))
        elapsed = time.perf_counter() - started
        _current, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()

        result = {
            "files": files,
            "vulns_per_file": vulns,
            "records": records_total,
            "level_mismatch": mismatches,
            "seconds": round(elapsed, 3),
            "seconds_per_file": round(elapsed / max(files, 1), 3),
            "peak_memory_mb": round(peak / 1024 / 1024, 1),
            "peak_per_file_mb": round(peak / 1024 / 1024 / max(files, 1), 2),
            "docx_disk_mb": round(gen_bytes / 1024 / 1024, 1),
            "images_disk_mb": round(_dir_size(images_dir) / 1024 / 1024, 2),
            "python": __import__("sys").version.split()[0],
        }
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    print("导入性能基准（顺序解析，与 worker 导入任务同路径）")
    print(f"  文件数 / 每份漏洞数：{result['files']} / {result['vulns_per_file']}")
    print(f"  解析记录数：{result['records']}（等级不一致 {result['level_mismatch']}）")
    print(f"  总耗时：{result['seconds']}s（{result['seconds_per_file']}s/份）")
    print(f"  tracemalloc 峰值：{result['peak_memory_mb']}MB（{result['peak_per_file_mb']}MB/份）")
    print(f"  磁盘：docx {result['docx_disk_mb']}MB + 解析图片 {result['images_disk_mb']}MB")
    print(f"  JSON：{out}")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Talos 导入性能基准（ROADMAP P1-5）")
    parser.add_argument("--files", type=int, default=20, help="报告份数（上限口径 20）")
    parser.add_argument("--vulns", type=int, default=12, help="每份报告的漏洞条数")
    parser.add_argument("--out", default="storage/bench_import.json", help="结果 JSON 路径")
    args = parser.parse_args()
    # 基准只做解析（不落库），但 settings 在非 DEBUG 下要求强密钥；先置位再导入 app
    import os

    os.environ.setdefault("VP_DEBUG", "1")
    run(args.files, args.vulns, Path(args.out))


if __name__ == "__main__":
    main()
