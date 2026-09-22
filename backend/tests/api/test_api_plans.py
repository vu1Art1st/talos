"""API 集成测试：工单（渗透 / 漏扫）。

工单（渗透/漏扫）：聚合筛选与嵌套条件、工单ID
搜索与序号占用、流程状态机、详情、统计与导出、人天与关联报告、Excel
导入、无漏洞工单完成口径。

原 tests/test_api.py（4626 行单体）按领域拆分而来：用例源码逐字未改，
仅按领域重新归位；共享夹具见 tests/conftest.py，共享 helper 见 tests/api/_helpers.py。
"""
import json
from io import BytesIO
import pytest
from httpx import AsyncClient

from _helpers import (
    _login_ready,
    _get_plan,
    _list_plan_names,
    _list_plans,
)

pytestmark = pytest.mark.asyncio(loop_scope="session")


async def test_testing_plan_filters(client: AsyncClient, auth: dict):
    """聚合筛选：多字段规则 + and/or 连接 + not 取反 + 派生/关联字段 + 非法入参。

    系统名/部门均使用唯一前缀，且断言限定在专属部门范围内，避免共享会话数据库
    中其他测试残留计划干扰精确集合比较。
    """
    DEPT = "筛选专用部门"
    A = "筛选专用系统-渗透A"
    B = "筛选专用系统-审计B"
    C = "筛选专用系统-渗透C"
    bodies = [
        {"system_name": A, "test_type": "渗透测试", "department": DEPT,
         "receive_time": "2026-01-01", "status": 10, "est_mandays": 3, "actual_mandays": 1.5},
        {"system_name": B, "test_type": "代码审计", "department": "筛选专用部门B",
         "receive_time": "2026-02-01", "status": 20, "est_mandays": 5},
        {"system_name": C, "test_type": "渗透测试", "department": DEPT,
         "receive_time": "2026-03-01", "status": 60, "est_mandays": 2},
    ]
    plans = []
    for body in bodies:
        resp = await client.post("/api/v1/testing-plans", headers=auth, json=body)
        assert resp.status_code == 200, resp.text
        plans.append(resp.json())
    try:
        def rule(field, op, value=None, *, not_=False, connector="and"):
            return {"field": field, "op": op, "value": value, "not": not_, "connector": connector}

        def q(*rules):
            return {"filters": json.dumps({"rules": list(rules)})}

        dept = rule("department", "eq", DEPT)

        # 单条件：部门等于
        assert await _list_plan_names(client, auth, q(dept)) == {A, C}
        # 文本包含
        assert await _list_plan_names(client, auth, q(rule("system_name", "contains", "筛选专用系统-渗透"))) == {A, C}
        # AND 组合：部门=专属部门 且 状态=10
        assert await _list_plan_names(client, auth, q(dept, rule("status", "eq", 10))) == {A}
        # NOT 取反：部门=专属部门 且 状态≠10
        assert await _list_plan_names(client, auth, q(dept, rule("status", "ne", 10))) == {C}
        # OR 组合：规则间按顺序左结合（AND 优先），用唯一系统名避免受其他测试数据影响
        assert await _list_plan_names(
            client, auth,
            q(rule("system_name", "eq", A), rule("system_name", "eq", C, connector="or")),
        ) == {A, C}
        # 数字比较与区间（限定专属部门）
        assert await _list_plan_names(client, auth, q(dept, rule("est_mandays", "gte", 3))) == {A}
        assert await _list_plan_names(client, auth, q(dept, rule("est_mandays", "between", [2, 3]))) == {A, C}
        # 实际人天：与预估人天同属 number 字段（前端字段清单缺口曾导致该列无法筛选）
        assert await _list_plan_names(client, auth, q(dept, rule("actual_mandays", "eq", 1.5))) == {A}
        assert await _list_plan_names(client, auth, q(dept, rule("actual_mandays", "between", [1, 2]))) == {A}
        # 未填实际人天的记录其值为 0，故 number 的「为空」等价于「NULL 或 0」
        assert await _list_plan_names(client, auth, q(dept, rule("actual_mandays", "is_empty"))) == {C}
        # 日期字符串上界（严格小于），空值被排除
        assert await _list_plan_names(client, auth, q(dept, rule("receive_time", "lt", "2026-02-01"))) == {A}
        # 空值筛选
        assert await _list_plan_names(client, auth, q(dept, rule("retest_notice_time", "is_empty"))) == {A, C}

        # 派生字段：自动生成工单ID精确匹配
        target = next(p for p in plans if p["system_name"] == A)
        assert target["ticket_id"]
        assert await _list_plan_names(client, auth, q(rule("ticket_id", "eq", target["ticket_id"]))) == {A}

        # 关联字段：认领后按测试人员筛选（限定专属部门）
        await client.post(f"/api/v1/testing-plans/{target['id']}/claim", headers=auth)
        assert await _list_plan_names(client, auth, q(dept, rule("testers", "contains", "admin"))) == {A}
        assert await _list_plan_names(client, auth, q(dept, rule("testers", "contains", "admin", not_=True))) == {C}

        # 非法字段 / 不匹配操作符返回 400
        resp = await client.get("/api/v1/testing-plans", headers=auth,
                                params=q(rule("not_exist", "eq", 1)))
        assert resp.status_code == 400
        resp = await client.get("/api/v1/testing-plans", headers=auth,
                                params=q(rule("system_name", "gt", 1)))
        assert resp.status_code == 400

        # stats 端点应用同一套筛选条件
        resp = await client.get("/api/v1/testing-plans/stats", headers=auth,
                                params=q(dept))
        assert resp.status_code == 200
        assert resp.json()["total_plans"] == 2
    finally:
        for p in plans:
            await client.delete(f"/api/v1/testing-plans/{p['id']}", headers=auth)

async def test_testing_plan_nested_filters(client: AsyncClient, auth: dict):
    """聚合筛选支持条件分组与嵌套：「需求晚于 X 且（状态 Y 或 Z）」、组级取反、三层嵌套与结构非法。

    断言限定在专属部门内，避免共享会话数据库中其他测试残留计划干扰集合比较。
    """
    DEPT = "嵌套筛选部门"
    A, B, C, D = (f"嵌套筛选系统-{s}" for s in "ABCD")
    bodies = [
        {"system_name": A, "test_type": "渗透测试", "department": DEPT,
         "receive_time": "2026-01-05", "status": 30},
        {"system_name": B, "test_type": "渗透测试", "department": DEPT,
         "receive_time": "2026-01-06", "status": 40},
        {"system_name": C, "test_type": "渗透测试", "department": DEPT,
         "receive_time": "2026-01-07", "status": 10},
        {"system_name": D, "test_type": "代码审计", "department": DEPT,
         "receive_time": "2026-01-08", "status": 30},
    ]
    plans = []
    for body in bodies:
        resp = await client.post("/api/v1/testing-plans", headers=auth, json=body)
        assert resp.status_code == 200, resp.text
        plans.append(resp.json())
    try:
        def rule(field, op, value=None, *, not_=False):
            return {"field": field, "op": op, "value": value, "not": not_}

        def cond(*children, logic="and", not_=False):
            return {"logic": logic, "not": not_, "children": list(children)}

        def q(node):
            return {"filters": json.dumps(node)}

        dept = rule("department", "eq", DEPT)

        # 「部门=专属 且（需求接收晚于 2026-01-05 且（状态=初测完成 或 状态=提请复测））」
        assert await _list_plan_names(client, auth, q(cond(
            dept,
            cond(rule("receive_time", "gt", "2026-01-05"),
                 cond(rule("status", "eq", 30), rule("status", "eq", 40), logic="or")),
        ))) == {B, D}

        # 组级取反：部门=专属 且 非（状态=未测试 或 测试类型=代码审计）
        assert await _list_plan_names(client, auth, q(cond(
            dept,
            cond(rule("status", "eq", 10), rule("test_type", "eq", "代码审计"),
                 logic="or", not_=True),
        ))) == {A, B}

        # 三层嵌套：部门=专属 且（状态=未测试 或（接收晚于 2026-01-07 且 状态=初测完成））
        assert await _list_plan_names(client, auth, q(cond(
            dept,
            cond(
                rule("status", "eq", 10),
                cond(rule("receive_time", "gt", "2026-01-07"), rule("status", "eq", 30)),
                logic="or",
            ),
        ))) == {C, D}

        # 空分组被忽略，不改变其它条件的语义
        assert await _list_plan_names(client, auth, q(cond(
            dept, cond(cond(), rule("system_name", "eq", A)),
        ))) == {A}

        # stats 端点应用同一套嵌套条件
        resp = await client.get("/api/v1/testing-plans/stats", headers=auth, params=q(cond(
            dept, cond(rule("status", "eq", 30), rule("status", "eq", 40), logic="or"),
        )))
        assert resp.status_code == 200
        assert resp.json()["total_plans"] == 3

        # 结构非法：非法逻辑运算符 / 嵌套过深 / 条件超量 → 400
        for bad in (
            cond(rule("status", "eq", 10), logic="xor"),
            cond(cond(cond(cond(cond(cond(rule("status", "eq", 10))))))),
            cond(*[rule("status", "eq", 10) for _ in range(51)]),
        ):
            resp = await client.get("/api/v1/testing-plans", headers=auth, params=q(bad))
            assert resp.status_code == 400, bad
    finally:
        for p in plans:
            await client.delete(f"/api/v1/testing-plans/{p['id']}", headers=auth)

async def test_testing_plan_search_by_ticket_id(client: AsyncClient, auth: dict):
    """关键词搜索支持工单ID：手动指定值 / 自动编号（完整 YYYYMMDD-N、日期段、序号段）。"""
    DEPT = "工单ID搜索部门"
    A = "工单ID搜索系统-手动"
    B = "工单ID搜索系统-自动"
    created = []

    # 自动编号：需求接收日期 2026-05-06 → 20260506-N
    resp = await client.post("/api/v1/testing-plans", headers=auth, json={
        "system_name": B, "test_type": "渗透测试", "department": DEPT,
        "receive_time": "2026-05-06",
    })
    assert resp.status_code == 200, resp.text
    auto = resp.json()
    created.append(auto)
    assert auto["ticket_id"].startswith("20260506-")

    # 手动指定工单ID（与自动编号同属一个搜索口径）
    manual_id = "ZDDD-2026-001"
    resp = await client.post("/api/v1/testing-plans", headers=auth, json={
        "system_name": A, "test_type": "渗透测试", "department": DEPT,
        "receive_time": "2026-05-06", "ticket_id_manual": manual_id,
    })
    assert resp.status_code == 200, resp.text
    manual = resp.json()
    created.append(manual)
    assert manual["ticket_id"] == manual_id

    def q(keyword: str) -> dict:
        # 限定专属部门，避免共享会话库中其他测试残留计划干扰集合比较
        return {"search": keyword, "department": DEPT}

    try:
        # 手动指定值（完整值与片段）
        assert await _list_plan_names(client, auth, q(manual_id)) == {A}
        assert await _list_plan_names(client, auth, q("ZDDD")) == {A}
        # 自动编号完整值 YYYYMMDD-N
        assert await _list_plan_names(client, auth, q(auto["ticket_id"])) == {B}
        # 自动编号的日期段（YYYYMMDD）与序号段
        assert await _list_plan_names(client, auth, q("20260506")) == {A, B}
        assert await _list_plan_names(client, auth, q(f"20260506-{auto['ticket_seq']}")) == {B}
        # 原有口径不受影响：系统名关键词仍可搜
        assert await _list_plan_names(client, auth, q("工单ID搜索系统")) == {A, B}
    finally:
        for p in created:
            await client.delete(f"/api/v1/testing-plans/{p['id']}", headers=auth)

async def test_testing_plan_workflow(client: AsyncClient, auth: dict):
    """测试计划工作台：认领/退出、录入漏洞统计重算、报告关联三方状态联动。"""
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={"system_name": "工作台系统", "test_type": "渗透测试", "department": "研发部"},
    )
    assert resp.status_code == 200, resp.text
    plan_id = resp.json()["id"]
    assert resp.json()["status"] == 10

    # 认领：当前用户加入测试人员，未测试自动进入初测中
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/claim", headers=auth)
    assert resp.status_code == 200, resp.text
    plan = resp.json()
    assert plan["status"] == 20
    assert "admin" in [u["username"] for u in plan["testers"]]

    # 认领幂等：重复调用不重复添加
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/claim", headers=auth)
    assert len(resp.json()["testers"]) == 1

    # 带计划录入漏洞：统计按等级自动重算
    resp = await client.post("/api/v1/assets", headers=auth, json={"name": "工作台资产"})
    asset_id = resp.json()["id"]
    resp = await client.post(
        "/api/v1/vulns/batch", headers=auth,
        json={
            "asset_ids": [asset_id],
            "vulns": [
                {"title": "计划漏洞A", "level": 20, "testing_plan_id": plan_id},
                {"title": "计划漏洞B", "level": 30, "testing_plan_id": plan_id},
            ],
        },
    )
    assert resp.status_code == 200, resp.text
    vul_a, vul_b = [v["id"] for v in resp.json()]
    plan = await _get_plan(client, auth, plan_id)
    assert (plan["stat_high"], plan["stat_medium"]) == (1, 1)
    assert {v["id"] for v in plan["vuls"]} == {vul_a, vul_b}

    # 不存在的计划被拒绝
    resp = await client.post(
        "/api/v1/vulns", headers=auth, json={"title": "坏计划", "testing_plan_id": 999999},
    )
    assert resp.status_code == 400

    # 生成报告并关联计划：报告带 testing_plan_id，计划进入初测完成
    resp = await client.post(
        "/api/v1/reports/from-vulns", headers=auth,
        json={"title": "工作台报告", "vul_ids": [vul_a, vul_b], "testing_plan_id": plan_id},
    )
    assert resp.status_code == 200, resp.text
    report_id = resp.json()["id"]
    assert resp.json()["testing_plan_id"] == plan_id
    plan = await _get_plan(client, auth, plan_id)
    assert plan["status"] == 30
    assert plan["first_test_done_time"]

    # 发起复测：计划进入复测中，漏洞进入复测中
    resp = await client.post(f"/api/v1/reports/{report_id}/retest", headers=auth)
    assert resp.status_code == 200, resp.text
    plan = await _get_plan(client, auth, plan_id)
    assert plan["status"] == 50
    for vid in (vul_a, vul_b):
        vul = (await client.get(f"/api/v1/vulns/{vid}", headers=auth)).json()
        assert vul["status"] == 55

    # 部分处理完：报告与计划均未完成
    resp = await client.post(
        f"/api/v1/vulns/{vul_a}/transition", headers=auth,
        json={"status": 60, "retest_html": "<p>复测通过</p>"},
    )
    assert resp.status_code == 200
    plan = await _get_plan(client, auth, plan_id)
    assert plan["status"] == 50

    # 全部已修复/已忽略：报告保持草稿（需求6），计划复测完成并记完成时间
    resp = await client.post(
        f"/api/v1/vulns/{vul_b}/transition", headers=auth, json={"status": 20},
    )
    assert resp.status_code == 200
    report = (await client.get(f"/api/v1/reports/{report_id}", headers=auth)).json()
    assert report["status"] == "draft"
    plan = await _get_plan(client, auth, plan_id)
    assert plan["status"] == 60
    assert plan["retest_done_time"]

    # 非管理员权限边界：状态修改与录入漏洞仅限认领者
    resp = await client.post(
        "/api/v1/roles", headers=auth,
        json={"name": "计划测试员", "permissions": ["special:manage", "vuln:submit"], "remark": ""},
    )
    role_id = resp.json()["id"]
    resp = await client.post(
        "/api/v1/users", headers=auth,
        json={"username": "plan_tester", "password": "Tester@123", "realname": "计划测试员",
              "email": "", "phone": "", "is_active": True, "role_id": role_id},
    )
    assert resp.status_code == 200, resp.text
    auth2 = await _login_ready(client, "plan_tester", "Tester@123")

    plan_body = {k: plan[k] for k in (
        "system_name", "test_type", "department", "receive_time", "first_test_done_time",
        "status", "retest_notice_time", "retest_done_time",
        "stat_critical", "stat_high", "stat_medium", "stat_low", "target_urls", "detail",
    )}

    # 未认领：修改状态 403，录入漏洞 403，但基本信息编辑放行
    resp = await client.put(
        f"/api/v1/testing-plans/{plan_id}", headers=auth2, json={**plan_body, "status": 40},
    )
    assert resp.status_code == 403
    resp = await client.post(
        "/api/v1/vulns", headers=auth2, json={"title": "越权录入", "testing_plan_id": plan_id},
    )
    assert resp.status_code == 403
    resp = await client.put(
        f"/api/v1/testing-plans/{plan_id}", headers=auth2,
        json={**plan_body, "department": "研发一部"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["department"] == "研发一部"

    # 认领后可修改状态；退出认领后再改状态回到 403
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/claim", headers=auth2)
    assert len(resp.json()["testers"]) == 2
    # 当前计划已处于复测完成(60)，认领后改回复测中(50)为合法流转
    resp = await client.put(
        f"/api/v1/testing-plans/{plan_id}", headers=auth2, json={**plan_body, "status": 50},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["status"] == 50
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/quit", headers=auth2)
    assert "plan_tester" not in [u["username"] for u in resp.json()["testers"]]
    # 退出认领后修改状态（60 为合法目标但无认领权限）回到 403
    resp = await client.put(
        f"/api/v1/testing-plans/{plan_id}", headers=auth2, json={**plan_body, "status": 60},
    )
    assert resp.status_code == 403

async def test_testing_plan_detail(client: AsyncClient, auth: dict):
    """单条计划详情端点：返回关联字段；不存在 404；无 special:manage 权限 403。"""
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={"system_name": "详情端点系统", "test_type": "渗透测试"},
    )
    plan_id = resp.json()["id"]
    resp = await client.get(f"/api/v1/testing-plans/{plan_id}", headers=auth)
    assert resp.status_code == 200, resp.text
    detail = resp.json()
    assert detail["id"] == plan_id
    for key in ("testers", "vuls", "reports", "retest_rounds", "retest_round_count"):
        assert key in detail

    resp = await client.get("/api/v1/testing-plans/999999", headers=auth)
    assert resp.status_code == 404

    # 无 special:manage 权限的用户被拒绝
    resp = await client.post(
        "/api/v1/roles", headers=auth,
        json={"name": "无专项权限", "permissions": ["vuln:submit"], "remark": ""},
    )
    role_id = resp.json()["id"]
    resp = await client.post(
        "/api/v1/users", headers=auth,
        json={"username": "no_special", "password": "Tester@123", "realname": "无专项",
              "email": "", "phone": "", "is_active": True, "role_id": role_id},
    )
    assert resp.status_code == 200, resp.text
    resp = await client.post(
        "/api/v1/auth/login", data={"username": "no_special", "password": "Tester@123"},
    )
    auth2 = {"Authorization": f"Bearer {resp.json()['access_token']}"}
    resp = await client.get(f"/api/v1/testing-plans/{plan_id}", headers=auth2)
    assert resp.status_code == 403

async def test_testing_plan_filter_stats_export(client: AsyncClient, auth: dict):
    """测试计划筛选（类型/部门/时间范围）、多维度统计与 Excel 双 sheet 导出。"""
    from openpyxl import load_workbook

    tag = "筛选统计专用部门"
    # 三条计划：不同状态 / 类型 / 接收时间，便于区分筛选与统计口径
    plans = [
        {"system_name": "过滤系统A", "department": tag, "test_type": "黑盒测试",
         "receive_time": "2026-01-15", "status": 20},   # 初测中
        {"system_name": "过滤系统B", "department": tag, "test_type": "白盒测试",
         "receive_time": "2026-03-20", "status": 60},   # 复测完成
        {"system_name": "过滤系统C", "department": tag, "test_type": "黑盒测试",
         "receive_time": "2026-06-10", "status": 10},   # 未测试
    ]
    for p in plans:
        resp = await client.post("/api/v1/testing-plans", headers=auth, json=p)
        assert resp.status_code == 200, resp.text

    # 部门 + 类型精确筛选
    resp = await client.get(
        "/api/v1/testing-plans", headers=auth,
        params={"department": tag, "test_type": "黑盒测试"},
    )
    names = {i["system_name"] for i in resp.json()["items"]}
    assert names == {"过滤系统A", "过滤系统C"}

    # 时间范围筛选（按需求接收时间）
    resp = await client.get(
        "/api/v1/testing-plans", headers=auth,
        params={"department": tag, "receive_from": "2026-02-01", "receive_to": "2026-04-01"},
    )
    names = {i["system_name"] for i in resp.json()["items"]}
    assert names == {"过滤系统B"}

    # 统计口径：本部门下 3 条计划
    resp = await client.get("/api/v1/testing-plans/stats", headers=auth, params={"department": tag})
    assert resp.status_code == 200, resp.text
    stats = resp.json()
    assert stats["total_plans"] == 3
    assert stats["retest_done_plans"] == 1          # 仅 B 为复测完成
    assert stats["first_test_count"] == 2           # A(20)+B(60) 达到初测中及以上
    assert stats["total_test_count"] == stats["first_test_count"] + stats["retest_count"]
    status_counts = {r["status"]: r["count"] for r in stats["by_status"]}
    assert status_counts == {10: 1, 20: 1, 60: 1}
    assert isinstance(stats["vulns_by_month"], list)

    # pending 待办流程筛选：仅 未测试(10)/初测中(20)/复测中(50)，复测完成(60) 被排除
    resp = await client.get(
        "/api/v1/testing-plans", headers=auth,
        params={"department": tag, "pending": True},
    )
    names = {i["system_name"] for i in resp.json()["items"]}
    assert names == {"过滤系统A", "过滤系统C"}
    assert all(i["status"] in (10, 20, 50) for i in resp.json()["items"])

    # pending 与单状态筛选同时传参时按 AND 处理：无交集则空结果
    resp = await client.get(
        "/api/v1/testing-plans", headers=auth,
        params={"department": tag, "pending": True, "status": 60},
    )
    assert resp.json()["total"] == 0

    # pending 统计口径与列表一致：本部门待办 2 条
    pending_stats = (await client.get(
        "/api/v1/testing-plans/stats", headers=auth, params={"department": tag, "pending": True},
    )).json()
    assert pending_stats["total_plans"] == 2
    assert pending_stats["retest_done_plans"] == 0

    # 空结果集：全 0，不报错
    empty = (await client.get(
        "/api/v1/testing-plans/stats", headers=auth, params={"department": "不存在的部门XYZ"},
    )).json()
    assert empty["total_plans"] == 0
    assert empty["total_test_count"] == 0

    # 导出：双 sheet，可被 openpyxl 回读
    resp = await client.get("/api/v1/testing-plans/export", headers=auth, params={"department": tag})
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"
    wb = load_workbook(BytesIO(resp.content))
    assert "渗透测试工单" in wb.sheetnames
    assert "统计汇总" in wb.sheetnames
    detail_rows = list(wb["渗透测试工单"].iter_rows(values_only=True))
    assert len(detail_rows) == 4  # 表头 + 3 行
    assert detail_rows[0][1] == "渗透测试工单名称"  # 表头第二列与 PLAN_EXCEL_HEADERS 一致

    # 导出同样支持 pending 筛选：仅 2 条待办计划
    resp = await client.get(
        "/api/v1/testing-plans/export", headers=auth,
        params={"department": tag, "pending": True},
    )
    assert resp.status_code == 200
    wb2 = load_workbook(BytesIO(resp.content))
    pending_rows = list(wb2["渗透测试工单"].iter_rows(values_only=True))
    assert len(pending_rows) == 3  # 表头 + 2 行
    assert "统计汇总" in wb2.sheetnames

async def test_testing_plan_mandays_and_reports(client: AsyncClient, auth: dict):
    """人天字段、人天统计（剩余预估人天=未测试计划预估人天之和）与计划反向展示关联报告。"""
    from openpyxl import load_workbook

    tag = "人天统计专用部门"
    plans = [
        {"system_name": "人天系统A", "department": tag, "status": 10,
         "est_mandays": 3, "actual_mandays": 0},     # 未测试 → 计入剩余
        {"system_name": "人天系统B", "department": tag, "status": 20,
         "est_mandays": 5, "actual_mandays": 2.5},
        {"system_name": "人天系统C", "department": tag, "status": 60,
         "est_mandays": 2, "actual_mandays": 4},
    ]
    plan_ids = []
    for p in plans:
        resp = await client.post("/api/v1/testing-plans", headers=auth, json=p)
        assert resp.status_code == 200, resp.text
        assert resp.json()["est_mandays"] == p["est_mandays"]
        assert resp.json()["actual_mandays"] == p["actual_mandays"]
        plan_ids.append(resp.json()["id"])

    # 人天统计：总预估/总实际/剩余预估（仅未测试状态）
    stats = (await client.get(
        "/api/v1/testing-plans/stats", headers=auth, params={"department": tag},
    )).json()
    assert stats["est_mandays_total"] == 10
    assert stats["actual_mandays_total"] == 6.5
    assert stats["remaining_est_mandays"] == 3

    # 导出明细含人天列，汇总含人天指标
    resp = await client.get("/api/v1/testing-plans/export", headers=auth, params={"department": tag})
    wb = load_workbook(BytesIO(resp.content))
    detail_rows = list(wb["渗透测试工单"].iter_rows(values_only=True))
    header = list(detail_rows[0])
    assert "预估人天" in header and "实际人天" in header
    est_col = header.index("预估人天")
    assert {r[est_col] for r in detail_rows[1:]} == {3, 5, 2}
    summary = {r[0]: r[1] for r in wb["统计汇总"].iter_rows(values_only=True) if r[0]}
    assert summary["预估人天总计"] == 10
    assert summary["实际人天总计"] == 6.5
    assert summary["剩余预估人天（未测试）"] == 3

    # 计划反向展示关联报告：录入漏洞→生成报告后 reports 列表可见
    plan_b = plan_ids[1]
    # 严格认领：需先认领计划才能建计划漏洞
    await client.post(f"/api/v1/testing-plans/{plan_b}/claim", headers=auth)
    resp = await client.post("/api/v1/assets", headers=auth, json={"name": "人天反向资产"})
    asset_id = resp.json()["id"]
    resp = await client.post(
        "/api/v1/vulns", headers=auth,
        json={"title": "人天反向漏洞", "level": 20, "asset_ids": [asset_id],
              "testing_plan_id": plan_b},
    )
    vul_id = resp.json()["id"]
    resp = await client.post(
        "/api/v1/reports/from-vulns", headers=auth,
        json={"title": "人天反向报告", "vul_ids": [vul_id], "testing_plan_id": plan_b},
    )
    assert resp.status_code == 200, resp.text
    report_id = resp.json()["id"]

    plan = await _get_plan(client, auth, plan_b)
    assert [r["id"] for r in plan["reports"]] == [report_id]
    assert plan["reports"][0]["title"] == "人天反向报告"
    assert plan["reports"][0]["status"] == "draft"

async def test_testing_plan_actual_mandays_override(client: AsyncClient, auth: dict):
    """实际人天修正：修正后不再被初测报告自动覆盖，取消修正后恢复自动计算。"""
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={"system_name": "修正人天系统", "department": "人天统计专用部门", "est_mandays": 2},
    )
    assert resp.status_code == 200, resp.text
    plan_id = resp.json()["id"]
    await client.post(f"/api/v1/testing-plans/{plan_id}/claim", headers=auth)
    asset_id = (await client.post(
        "/api/v1/assets", headers=auth, json={"name": "修正人天资产"},
    )).json()["id"]
    vul_id = (await client.post(
        "/api/v1/vulns", headers=auth,
        json={"title": "修正人天漏洞", "level": 20, "asset_ids": [asset_id],
              "testing_plan_id": plan_id},
    )).json()["id"]
    resp = await client.post(
        "/api/v1/reports/from-vulns", headers=auth,
        json={"title": "修正人天报告", "vul_ids": [vul_id], "testing_plan_id": plan_id},
    )
    assert resp.status_code == 200, resp.text
    report = resp.json()
    report_id = report["id"]

    # 给初测报告填写测试周期：3 天 → 计划实际人天自动 = 3
    save_body = {
        "title": report["title"], "project_name": report["project_name"],
        "author": report["author"], "revision": report["revision"],
        "test_start": "2026-08-01", "test_end": "2026-08-03",
        "sections": report["sections"],
    }
    resp = await client.put(f"/api/v1/reports/{report_id}", headers=auth, json=save_body)
    assert resp.status_code == 200, resp.text
    report = resp.json()
    assert report["actual_mandays"] == 3
    plan = await _get_plan(client, auth, plan_id)
    assert plan["actual_mandays"] == 3
    assert plan["actual_mandays_override"] is False

    # 修正：手动输入 9.5，保存后不再被初测报告自动覆盖
    body = plan
    for k in ("id", "ticket_seq", "ticket_id", "testers", "vuls", "reports",
              "retest_rounds", "retest_round_count", "create_time", "update_time"):
        body.pop(k, None)
    body["actual_mandays"] = 9.5
    body["actual_mandays_override"] = True
    resp = await client.put(f"/api/v1/testing-plans/{plan_id}", headers=auth, json=body)
    assert resp.status_code == 200, resp.text
    assert resp.json()["actual_mandays"] == 9.5
    assert resp.json()["actual_mandays_override"] is True

    # 初测报告测试周期变更为 5 天，refresh_mandays 因修正标志跳过，修正值保持不变
    save_body["revision"] = report["revision"]
    save_body["test_end"] = "2026-08-05"
    resp = await client.put(f"/api/v1/reports/{report_id}", headers=auth, json=save_body)
    assert resp.status_code == 200, resp.text
    plan = await _get_plan(client, auth, plan_id)
    assert plan["actual_mandays"] == 9.5

    # 取消修正：恢复为初测报告计算的 5 天
    body["actual_mandays"] = 99  # 该值会被自动计算覆盖
    body["actual_mandays_override"] = False
    resp = await client.put(f"/api/v1/testing-plans/{plan_id}", headers=auth, json=body)
    assert resp.status_code == 200, resp.text
    plan = resp.json()
    assert plan["actual_mandays_override"] is False
    assert plan["actual_mandays"] == 5

async def test_testing_plan_excel_import(client: AsyncClient, auth: dict):
    """测试计划 Excel 导入：模板下载、按 ID 更新、无 ID 新增、非法行报错。"""
    from openpyxl import Workbook

    # 模板下载
    resp = await client.get("/api/v1/testing-plans/import/template", headers=auth)
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"

    # 先建一条计划供导入更新
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={"system_name": "导入前系统", "department": "导入专用部门", "est_mandays": 1},
    )
    exist_id = resp.json()["id"]

    # 表头与模板一致（PLAN_EXCEL_HEADERS 20 列，含「渗透测试工单名称」列）
    headers_row = ["ID", "渗透测试工单名称", "测试系统", "测试类型", "所属部门", "工单ID",
                   "工单提起时间", "状态", "测试人员", "需求接收", "初测完成", "复测通知",
                   "复测完成", "预估人天", "实际人天", "超危数", "高危数", "中危数",
                   "低危数", "复测轮数"]
    wb = Workbook()
    ws = wb.active
    ws.append(headers_row)
    # 无 ID → 新增（测试人员按用户名匹配 admin）
    ws.append(["", "", "导入新增系统", "渗透测试", "导入专用部门", "", "",
               "初测中", "admin", "2025-11-01", "2025-11-05", "", "", 3.5, 1, 1, 2, 0, 0, 0])
    # 有 ID → 更新同一条计划
    ws.append([exist_id, "", "导入后系统", "白盒测试", "导入专用部门", "", "",
               "复测完成", "", "2025-10-01", "", "", "2025-12-31", 6, 5.5, 0, 0, 0, 0, 0])
    # 缺测试系统 → 失败
    ws.append(["", "", "", "黑盒测试", "导入专用部门", "", "", "", "", "", "", "", "",
               0, 0, 0, 0, 0, 0, 0])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = await client.post(
        "/api/v1/testing-plans/import", headers=auth,
        files={"file": ("plans.xlsx", buf,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    result = resp.json()
    assert result["total"] == 3
    assert result["created"] == 1
    assert result["updated"] == 1
    assert result["failed"] == 1
    assert "测试系统为必填项" in result["errors"][0]

    # 新增行校验：状态/人天/测试人员匹配
    resp = await client.get(
        "/api/v1/testing-plans", headers=auth, params={"search": "导入新增系统"},
    )
    created = resp.json()["items"][0]
    assert created["status"] == 20
    assert created["est_mandays"] == 3.5
    assert created["stat_critical"] == 1 and created["stat_high"] == 2
    assert "admin" in [u["username"] for u in created["testers"]]

    # 更新行校验：同 ID 字段被覆盖
    updated = await _get_plan(client, auth, exist_id)
    assert updated["system_name"] == "导入后系统"
    assert updated["status"] == 60
    assert updated["est_mandays"] == 6
    assert updated["actual_mandays"] == 5.5

    # 非 xlsx 拒绝
    resp = await client.post(
        "/api/v1/testing-plans/import", headers=auth,
        files={"file": ("bad.txt", b"hello", "text/plain")},
    )
    assert resp.status_code == 400

async def test_ticket_seq_increment_no_reuse(client: AsyncClient, auth: dict):
    """工单ID自动分配采用「当日最大编号+1」：删除/释放的编号不复用，仅手动可选用。"""
    body = {"department": "递增部门", "receive_time": "2026-01-01"}

    resp = await client.post(
        "/api/v1/testing-plans", headers=auth, json={**body, "system_name": "递增系统A"},
    )
    assert resp.status_code == 200, resp.text
    plan_a = resp.json()
    assert plan_a["ticket_id"] == "20260101-1"

    resp = await client.post(
        "/api/v1/testing-plans", headers=auth, json={**body, "system_name": "递增系统B"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["ticket_id"] == "20260101-2"

    # 删除占用 20260101-1 的计划后，自动分配不复用空洞，继续递增到 3
    resp = await client.delete(f"/api/v1/testing-plans/{plan_a['id']}", headers=auth)
    assert resp.status_code == 200, resp.text
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth, json={**body, "system_name": "递增系统C"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["ticket_id"] == "20260101-3"

    # 被释放的 20260101-1 仍可手动指定给新工单
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={**body, "system_name": "递增系统D", "ticket_id_manual": "20260101-1"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["ticket_id"] == "20260101-1"

async def test_ticket_id_manual_occupancy_and_ghost(client: AsyncClient, auth: dict):
    """工单编号占用口径统一：自动分配跳过 manual 占用；manual 记录底层 seq 不产生幽灵占用。"""
    body = {"department": "占用部门", "receive_time": "2026-07-30"}

    # A 手动指定 20260730-2（其 ticket_seq=0，不参与自动分配）
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={**body, "system_name": "占用系统A", "ticket_id_manual": "20260730-2"},
    )
    assert resp.status_code == 200, resp.text

    # 自动工单 B：max(纯自动 seq, manual N)=2 -> 新序号 3，编号 20260730-3（跳过 A 占用的 2）
    resp = await client.post("/api/v1/testing-plans", headers=auth, json={**body, "system_name": "占用系统B"})
    assert resp.status_code == 200, resp.text
    assert resp.json()["ticket_id"] == "20260730-3"

    # 自动工单 C：max(3, 2)+1=4 -> 20260730-4（单调递增）
    resp = await client.post("/api/v1/testing-plans", headers=auth, json={**body, "system_name": "占用系统C"})
    assert resp.status_code == 200, resp.text
    assert resp.json()["ticket_id"] == "20260730-4"

    # 幽灵占用场景：D 自动得到 20260730-5 后手动改为 20260730-88（底层 seq 保留 5）
    resp = await client.post("/api/v1/testing-plans", headers=auth, json={**body, "system_name": "占用系统D"})
    plan_d = resp.json()
    assert plan_d["ticket_id"] == "20260730-5"
    resp = await client.put(
        f"/api/v1/testing-plans/{plan_d['id']}", headers=auth,
        json={**body, "system_name": "占用系统D", "ticket_id_manual": "20260730-88"},
    )
    assert resp.status_code == 200, resp.text

    # E 手动指定 20260730-5（D 显示为 20260730-88，无任何工单显示 5）→ 应放行
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={**body, "system_name": "占用系统E", "ticket_id_manual": "20260730-5"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["ticket_id"] == "20260730-5"

async def test_plan_search_ticket_id_excludes_ghost_seq(client: AsyncClient, auth: dict):
    """幽灵序号不得被搜索 / 聚合筛选误命中（2026-09-22 缺陷回归）。

    现场：某记录先自动得到 20260730-1，手工改号后底层 ticket_seq 仍为 1；搜索 20260730-1 时
    该记录与真正显示 20260730-1 的记录一起命中 → 一次搜索返回两条。口径与
    `ticket_service.check_ticket_id_unique` 对齐：手动指定了编号的记录，其残留 ticket_seq 不参与命中。
    """
    DEPT = "幽灵序号搜索部门"
    base = {"test_type": "渗透测试", "department": DEPT, "receive_time": "2026-07-30"}
    created: list[int] = []

    resp = await client.post("/api/v1/testing-plans", headers=auth,
                             json={**base, "system_name": "幽灵序号系统-自动"})
    assert resp.status_code == 200, resp.text
    ghost = resp.json()
    created.append(ghost["id"])
    ghost_ticket = ghost["ticket_id"]  # 自动编号（此刻尚无任何记录显示它）
    assert ghost_ticket.startswith("20260730-")

    # 手工改号为另一个日期的编号：底层 ticket_seq 残留成「幽灵序号」
    resp = await client.put(f"/api/v1/testing-plans/{ghost['id']}", headers=auth,
                            json={**base, "system_name": "幽灵序号系统-自动",
                                  "ticket_id_manual": "20991231-99"})
    assert resp.status_code == 200, resp.text
    assert resp.json()["ticket_id"] == "20991231-99"

    # 另一条记录接管该显示编号
    resp = await client.post("/api/v1/testing-plans", headers=auth,
                             json={**base, "system_name": "幽灵序号系统-接管",
                                   "ticket_id_manual": ghost_ticket})
    assert resp.status_code == 200, resp.text
    takeover_id = resp.json()["id"]
    created.append(takeover_id)

    def assert_only_takeover(items: list[dict]) -> None:
        ids = [p["id"] for p in items]
        assert ghost["id"] not in ids, f"幽灵序号记录被误命中：{items}"
        assert takeover_id in ids, f"接管记录未被命中：{items}"

    try:
        # 关键词搜索（修复前会同时返回幽灵记录）
        assert_only_takeover(await _list_plans(
            client, auth, {"search": ghost_ticket, "department": DEPT},
        ))

        # 聚合筛选同口径（plan_query._ticket_id_filter_expr）
        filters = json.dumps({"logic": "and", "not": False, "children": [
            {"kind": "rule", "field": "department", "op": "eq", "value": DEPT},
            {"kind": "rule", "field": "ticket_id", "op": "eq", "value": ghost_ticket},
        ]})
        assert_only_takeover(await _list_plans(client, auth, {"filters": filters}))
    finally:
        for plan_id in created:
            await client.delete(f"/api/v1/testing-plans/{plan_id}", headers=auth)

async def test_plan_complete_no_vuln_flow(client: AsyncClient, auth: dict):
    """无漏洞闭环：无漏洞完结 → 测试通过 + 无漏洞报告 → 重复确认拒绝 → 补录漏洞自动重开。"""
    # meta 字典包含新增的「测试通过」状态
    meta = (await client.get("/api/v1/meta", headers=auth)).json()
    assert meta["testing_plan_status"]["70"] == "测试通过"

    # 初测中且无漏洞的计划
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={"system_name": "无漏洞闭环系统", "test_type": "渗透测试",
              "receive_time": "2026-08-01", "status": 20},
    )
    assert resp.status_code == 200, resp.text
    plan_id = resp.json()["id"]

    # 确认无漏洞完结：状态流转、初测完成打点、结论记录、无漏洞报告生成
    resp = await client.post(
        f"/api/v1/testing-plans/{plan_id}/complete-no-vuln", headers=auth,
        json={"conclusion": "覆盖 OWASP Top 10 主要攻击面，未发现安全漏洞",
              "generate_report": True, "title": ""},
    )
    assert resp.status_code == 200, resp.text
    plan = resp.json()
    assert plan["status"] == 70
    assert plan["no_vul_conclusion"] == "覆盖 OWASP Top 10 主要攻击面，未发现安全漏洞"
    assert plan["first_test_done_time"]
    assert len(plan["reports"]) == 1
    report_brief = plan["reports"][0]
    assert "无漏洞闭环系统" in report_brief["title"]
    assert "渗透测试报告（无漏洞）" in report_brief["title"]
    # 无漏洞报告按初测报告口径计入计划实际人天（接收日期至确认当天）
    assert plan["actual_mandays"] >= 1

    # 报告单章节「测试结论」，无漏洞关联，结论含安全测试通过文案与补充说明
    resp = await client.get(f"/api/v1/reports/{report_brief['id']}", headers=auth)
    assert resp.status_code == 200, resp.text
    detail = resp.json()
    assert detail["testing_plan_id"] == plan_id
    assert len(detail["sections"]) == 1
    section = detail["sections"][0]
    assert section["title"] == "测试结论"
    assert section["vul_id"] is None
    assert "未发现" in section["content_html"]
    assert "安全测试通过" in section["content_html"]
    assert "OWASP" in section["content_html"]

    # 重复确认被拒绝
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/complete-no-vuln", headers=auth, json={})
    assert resp.status_code == 400

    # 补录漏洞后计划自动重开为「初测中」（关联需先认领）
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/claim", headers=auth)
    assert resp.status_code == 200, resp.text
    resp = await client.post(
        "/api/v1/vulns", headers=auth, json={"title": "无漏洞闭环补录漏洞", "level": 30},
    )
    assert resp.status_code == 200, resp.text
    vul_id = resp.json()["id"]
    resp = await client.post(
        f"/api/v1/testing-plans/{plan_id}/attach-vulns", headers=auth, json={"vul_ids": [vul_id]},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["status"] == 20
    # 重开后不可再确认无漏洞（存在关联漏洞），无漏洞测试结论保留以便追溯
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/complete-no-vuln", headers=auth, json={})
    assert resp.status_code == 400
    plan = (await client.get(f"/api/v1/testing-plans/{plan_id}", headers=auth)).json()
    assert plan["no_vul_conclusion"] == "覆盖 OWASP Top 10 主要攻击面，未发现安全漏洞"

async def test_plan_complete_no_vuln_requires_no_vulns(client: AsyncClient, auth: dict):
    """存在关联漏洞的计划不能确认无漏洞；直接创建漏洞到已通过计划同样触发重开。"""
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={"system_name": "有漏洞系统", "receive_time": "2026-08-01", "status": 20},
    )
    assert resp.status_code == 200, resp.text
    plan_id = resp.json()["id"]
    # 录入/关联漏洞需先认领计划
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/claim", headers=auth)
    assert resp.status_code == 200, resp.text
    resp = await client.post(
        "/api/v1/vulns", headers=auth,
        json={"title": "有漏洞系统-漏洞1", "level": 40, "testing_plan_id": plan_id},
    )
    assert resp.status_code == 200, resp.text
    vid = resp.json()["id"]

    # 有关联漏洞时拒绝无漏洞完结
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/complete-no-vuln", headers=auth, json={})
    assert resp.status_code == 400

    # 移除漏洞后不生成报告仅流转状态与记录结论
    resp = await client.delete(f"/api/v1/vulns/{vid}", headers=auth)
    assert resp.status_code == 200
    resp = await client.post(
        f"/api/v1/testing-plans/{plan_id}/complete-no-vuln", headers=auth,
        json={"conclusion": "仅完结不生成报告", "generate_report": False},
    )
    assert resp.status_code == 200, resp.text
    plan = resp.json()
    assert plan["status"] == 70
    assert plan["reports"] == []

    # 直接向已通过计划创建漏洞（不经 attach-vulns）同样自动重开
    resp = await client.post(
        "/api/v1/vulns", headers=auth,
        json={"title": "有漏洞系统-漏洞2", "level": 40, "testing_plan_id": plan_id},
    )
    assert resp.status_code == 200, resp.text
    plan = (await client.get(f"/api/v1/testing-plans/{plan_id}", headers=auth)).json()
    assert plan["status"] == 20
