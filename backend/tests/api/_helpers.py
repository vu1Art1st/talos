"""tests/api 的共享 helper 与数据常量（由原 tests/test_api.py 逐字搬移）。

这些函数贯穿多个领域模块，故独立成本模块；命名以 `_` 开头，pytest 不会收集。
"""
import asyncio
from io import BytesIO
from httpx import AsyncClient


_SEARCH_ENTRIES = [
    {"vulnerability_name": "搜索用例-SQL注入（CVE-2099-1111）", "vul_type": 10,
     "severity_level": 10, "description_html": "<p>搜索用例描述：拼接 SQL 语句导致注入。</p>"},
    {"vulnerability_name": "搜索用例-Shiro认证绕过（CVE-2099-2222）", "vul_type": 40,
     "severity_level": 20, "description_html": "<p>搜索用例描述：路径规范化不一致。</p>"},
    {"vulnerability_name": "搜索用例-服务器弱口令", "vul_type": 65, "severity_level": 30,
     "description_html": "<p>搜索用例描述：默认口令未修改。</p>",
     "references": ["https://example.com/uniq-ref-3333"]},
]


async def _wait_batch(client: AsyncClient, auth: dict, batch_id: int) -> dict:
    for _ in range(50):
        resp = await client.get(f"/api/v1/imports/{batch_id}", headers=auth)
        assert resp.status_code == 200
        detail = resp.json()
        if detail["batch"]["status"] in ("parsed", "failed", "confirmed"):
            return detail
        await asyncio.sleep(0.2)
    raise AssertionError("导入批次解析超时")

def _build_report_docx(system_name: str, target_url: str, target_ip: str, sections: list[tuple[str, str, str, bool]],
                       status_texts: list[str] | None = None,
                       testers: list[str] | None = None, test_start: str = "", test_end: str = "",
                       test_account: str = ""):
    """构造平台报告格式 docx：测试目标表 + 时间与人员表（可选）+ 风险问题汇总表 + 风险问题详情章节。

    sections: [(标题, 等级文本, 类型文本, 是否已修复)]
    status_texts: 可选，逐条覆盖汇总表「修复状态」文本（如「部分未修复」），默认已修复/未修复。
    testers/test_start/test_end/test_account: 可选，构造「时间与人员」表与测试账号，
        用于验证参测人员/测试周期/测试账号解析回填。
    """
    from docx import Document

    doc = Document()
    # 测试目标表
    t = doc.add_table(rows=5, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "业务系统名称", system_name
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "被测系统URL", target_url
    t.rows[2].cells[0].text, t.rows[2].cells[1].text = "被测系统域名", ""
    t.rows[3].cells[0].text, t.rows[3].cells[1].text = "被测系统IP", target_ip
    t.rows[4].cells[0].text, t.rows[4].cells[1].text = "被测测试账号", test_account
    # 时间与人员表（结构与解析器 _parse_schedule_table 匹配：表头在第4行，参测人员从第5行起）
    if testers:
        sche = doc.add_table(rows=4 + len(testers), cols=4)
        sche.rows[0].cells[0].text = "测试工作时间段"
        sche.rows[1].cells[0].text, sche.rows[1].cells[1].text = "起始时间", test_start
        sche.rows[1].cells[2].text, sche.rows[1].cells[3].text = "结束时间", test_end
        for i, h in enumerate(("参测人员", "所属部门", "人员角色", "人员分工")):
            sche.rows[3].cells[i].text = h
        for i, name in enumerate(testers):
            sche.rows[4 + i].cells[0].text = name
    # 风险问题汇总表
    doc.add_heading("风险问题汇总", level=1)
    s = doc.add_table(rows=1 + len(sections), cols=4)
    for i, h in enumerate(("问题等级", "风险类型", "风险问题", "修复状态")):
        s.rows[0].cells[i].text = h
    for ri, (title, level_text, type_text, fixed) in enumerate(sections, start=1):
        s.rows[ri].cells[0].text = level_text
        s.rows[ri].cells[1].text = type_text
        s.rows[ri].cells[2].text = title
        if status_texts is not None and ri - 1 < len(status_texts):
            s.rows[ri].cells[3].text = status_texts[ri - 1]
        else:
            s.rows[ri].cells[3].text = "已修复" if fixed else "未修复"
    # 风险问题详情
    doc.add_heading("风险问题详情", level=1)
    for title, _lvl, _typ, fixed in sections:
        suffix = "（已修复）" if fixed else "（未修复）"
        doc.add_heading(f"{title}{suffix}", level=3)
        doc.add_paragraph("漏洞链接")
        doc.add_paragraph(target_url)
        doc.add_paragraph("漏洞描述")
        doc.add_paragraph("此处为漏洞描述内容。")
        doc.add_paragraph("漏洞证明")
        doc.add_paragraph("此处为漏洞证明内容。")
        doc.add_paragraph("修复建议")
        doc.add_paragraph("此处为修复建议内容。")
        doc.add_paragraph("20260728漏洞复测")
        doc.add_paragraph("复测详情：已按建议整改。")
    return doc

async def _import_report(client: AsyncClient, auth: dict, filename: str, doc):
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    resp = await client.post(
        "/api/v1/imports", headers=auth,
        files={"file": (filename, buf,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, resp.text
    batch_id = resp.json()["id"]
    detail = await _wait_batch(client, auth, batch_id)
    assert detail["batch"]["status"] == "parsed", detail
    assert detail["batch"]["doc_kind"] == "report", detail
    rec_ids = [r["id"] for r in detail["records"]]
    resp = await client.post(
        f"/api/v1/imports/{batch_id}/confirm", headers=auth,
        json={"record_ids": rec_ids},
    )
    assert resp.status_code == 200, resp.text
    return detail["records"], resp.json()

async def _upload_report_batch(client: AsyncClient, auth: dict, system_name: str, filename: str) -> int:
    """上传报告格式 docx 并等待解析完成（不确认入库），返回 batch_id。"""
    doc = _build_report_docx(
        system_name, f"http://10.9.9.9/{system_name}", "10.9.9.9",
        sections=[(f"{system_name}漏洞A", "高危", "逻辑漏洞", False)],
    )
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    resp = await client.post(
        "/api/v1/imports", headers=auth,
        files={"file": (filename, buf,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, resp.text
    batch_id = resp.json()["id"]
    detail = await _wait_batch(client, auth, batch_id)
    assert detail["batch"]["status"] == "parsed", detail
    assert detail["batch"]["doc_kind"] == "report", detail
    return batch_id

async def _find_vuln(client: AsyncClient, auth: dict, keyword: str) -> dict:
    resp = await client.get("/api/v1/vulns", headers=auth, params={"search": keyword})
    items = resp.json()["items"]
    assert items, f"未找到漏洞：{keyword}"
    return items[0]

async def _login_ready(client: AsyncClient, username: str, password: str) -> dict:
    """登录新建账号并完成首次改密，返回可用请求头。

    新建用户默认 `must_change_password=True`，服务端会拦截除 `/auth/*` 之外的一切接口
    （安全审计 批次 E-6：前端弹框强制的服务端兜底），故测试须先改密再调用业务接口。
    """
    resp = await client.post("/api/v1/auth/login", data={"username": username, "password": password})
    assert resp.status_code == 200, resp.text
    headers = {"Authorization": f"Bearer {resp.json()['access_token']}"}
    resp = await client.post(
        "/api/v1/auth/password", headers=headers,
        json={"old_password": password, "new_password": f"{password}x"},
    )
    assert resp.status_code == 200, resp.text
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}

async def _user_with_perms(client: AsyncClient, auth: dict, username: str, perms: list[str]) -> dict:
    """建「角色 + 用户」并返回可用请求头（角色名 = 用户名 + 角色，便于用例间隔离）。"""
    resp = await client.post(
        "/api/v1/roles", headers=auth,
        json={"name": f"{username}角色", "permissions": perms, "remark": ""},
    )
    assert resp.status_code == 200, resp.text
    resp = await client.post(
        "/api/v1/users", headers=auth,
        json={"username": username, "password": "Tester@123", "realname": username,
              "email": "", "phone": "", "is_active": True, "role_id": resp.json()["id"]},
    )
    assert resp.status_code == 200, resp.text
    return await _login_ready(client, username, "Tester@123")

async def _wait_job(client: AsyncClient, auth: dict, report_id: int, job_id: int) -> dict:
    for _ in range(50):
        resp = await client.get(f"/api/v1/reports/{report_id}/exports", headers=auth)
        job = next(j for j in resp.json() if j["id"] == job_id)
        if job["status"] in ("done", "failed"):
            return job
        await asyncio.sleep(0.2)
    raise AssertionError("导出任务超时")

async def _get_plan(client: AsyncClient, auth: dict, plan_id: int) -> dict:
    resp = await client.get("/api/v1/testing-plans", headers=auth, params={"size": 100})
    return next(p for p in resp.json()["items"] if p["id"] == plan_id)

async def _list_plan_names(client: AsyncClient, auth: dict, params: dict) -> set[str]:
    resp = await client.get("/api/v1/testing-plans", headers=auth, params={"size": 100, **params})
    assert resp.status_code == 200, resp.text
    return {p["system_name"] for p in resp.json()["items"]}

async def _seed_search_entries(client: AsyncClient, auth: dict) -> list[dict]:
    created = []
    for e in _SEARCH_ENTRIES:
        resp = await client.post("/api/v1/knowledge", headers=auth, json=e)
        assert resp.status_code == 200, resp.text
        created.append(resp.json())
    return created

def _many_urls(count: int, prefix: str = "https://many.example.com/api") -> list[str]:
    return [f"{prefix}/{i}" for i in range(count)]
