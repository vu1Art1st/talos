"""API 集成测试：Word 报告导入与导入批次。

Word 报告解析入库与导入批次：章节顺序、等级不一致提醒、批量确认（含时序与去
重）、自动导出、作者与公开 URL 解析、知识库回填。

原 tests/test_api.py（4626 行单体）按领域拆分而来：用例源码逐字未改，
仅按领域重新归位；共享夹具见 tests/conftest.py，共享 helper 见 tests/api/_helpers.py。
"""
from io import BytesIO
import pytest
from httpx import AsyncClient

from _helpers import (
    _wait_batch,
    _build_report_docx,
    _import_report,
    _upload_report_batch,
    _find_vuln,
)

pytestmark = pytest.mark.asyncio(loop_scope="session")


async def test_word_import_flow(client: AsyncClient, auth: dict):
    # 下载官方模板作为上传样例
    resp = await client.get("/api/v1/imports/template", headers=auth)
    assert resp.status_code == 200
    docx_bytes = resp.content

    # 非 docx 拒绝
    resp = await client.post(
        "/api/v1/imports", headers=auth,
        files={"file": ("bad.txt", b"hello", "text/plain")},
    )
    assert resp.status_code == 400

    # 上传并等待后台解析（DISABLE_QUEUE 下进程内执行）
    resp = await client.post(
        "/api/v1/imports", headers=auth,
        files={"file": ("样例报告.docx", BytesIO(docx_bytes),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, resp.text
    batch_id = resp.json()["id"]

    detail = await _wait_batch(client, auth, batch_id)
    assert detail["batch"]["status"] == "parsed", detail
    records = detail["records"]
    assert len(records) == 1
    rec = records[0]
    assert rec["status"] == "parsed"
    assert rec["level"] == 20

    # 预览修正标题后确认入库
    resp = await client.put(
        f"/api/v1/imports/records/{rec['id']}", headers=auth,
        json={"title": "后台登录接口存在SQL注入(已确认)"},
    )
    assert resp.status_code == 200

    resp = await client.post(
        f"/api/v1/imports/{batch_id}/confirm", headers=auth,
        json={"record_ids": [rec["id"]]},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["created"] == 1

    # 入库漏洞 source=0（来源未选择；Word导入不再单列来源）
    resp = await client.get(
        "/api/v1/vulns", headers=auth, params={"search": "已确认"}
    )
    items = resp.json()["items"]
    assert len(items) == 1
    assert items[0]["source"] == 0

async def test_report_import_partial_fixed_flow(client: AsyncClient, auth: dict):
    """复测报告含未修复漏洞：计划应为复测中(50)、未修复漏洞置「复测未修复」(50+is_retest)、自动建资产与报告。"""
    system_name = "综合办公系统ZZ"
    target_url = "http://10.9.9.9/officezz"
    doc = _build_report_docx(
        system_name, target_url, "10.9.9.9",
        sections=[
            ("平行越权访问漏洞ZZ", "高危", "逻辑漏洞", True),
            ("敏感信息泄露漏洞ZZ", "中危", "信息泄露", False),
        ],
    )
    _records, result = await _import_report(
        client, auth, "20260728综合办公系统ZZ渗透测试复测报告.docx", doc,
    )
    assert result["created"] == 2

    # 计划：复测中(50)，不是复测完成(60)
    resp = await client.get("/api/v1/testing-plans", headers=auth, params={"search": system_name})
    plans = resp.json()["items"]
    assert len(plans) == 1, plans
    assert plans[0]["status"] == 50

    # 已修复漏洞 → 60；未修复漏洞 → 50 且 is_retest=True（展示层为「复测未修复/复测未通过」）
    fixed_vuln = await _find_vuln(client, auth, "平行越权访问漏洞ZZ")
    assert fixed_vuln["status"] == 60
    unfixed_vuln = await _find_vuln(client, auth, "敏感信息泄露漏洞ZZ")
    assert unfixed_vuln["status"] == 50
    assert unfixed_vuln["is_retest"] is True
    assert any(a["name"] == system_name for a in unfixed_vuln["assets"])

    # 自动创建资产：系统名匹配，被测 URL 入内网地址
    resp = await client.get("/api/v1/assets", headers=auth, params={"search": system_name})
    assets = [a for a in resp.json()["items"] if a["name"] == system_name]
    assert assets, "未自动创建资产"
    assert target_url in (assets[0]["internal_urls"] or [])

    # 自动创建报告：显示在报告中心，草稿态，章节数与漏洞数一致
    resp = await client.get("/api/v1/reports", headers=auth, params={"search": system_name})
    reports = [r for r in resp.json()["items"] if r["project_name"] == system_name]
    assert reports, "未自动创建报告"
    report = reports[0]
    assert report["status"] == "draft"
    assert report["testing_plan_id"] == plans[0]["id"]
    detail = await client.get(f"/api/v1/reports/{report['id']}", headers=auth)
    assert len(detail.json()["sections"]) == 2

async def test_report_import_all_fixed_flow(client: AsyncClient, auth: dict):
    """复测报告全部修复：计划复测完成(60)、报告保持草稿(draft，定稿由导出 Word 驱动)。"""
    system_name = "门户系统ZZ"
    doc = _build_report_docx(
        system_name, "http://10.8.8.8/portalzz", "10.8.8.8",
        sections=[("命令执行漏洞ZZ", "高危", "命令执行漏洞", True)],
    )
    _records, result = await _import_report(
        client, auth, "20260728门户系统ZZ渗透测试复测报告.docx", doc,
    )
    assert result["created"] == 1

    resp = await client.get("/api/v1/testing-plans", headers=auth, params={"search": system_name})
    plans = resp.json()["items"]
    assert len(plans) == 1 and plans[0]["status"] == 60

    resp = await client.get("/api/v1/reports", headers=auth, params={"search": system_name})
    reports = [r for r in resp.json()["items"] if r["project_name"] == system_name]
    assert reports and reports[0]["status"] == "draft"

async def test_report_import_all_fixed_keeps_plan_when_other_vuln_open(
    client: AsyncClient, auth: dict,
):
    """复测报告全部修复，但工单仍存在其他未闭环漏洞（未纳入该报告）：工单保持「复测中」。

    回归：导入路径早期以「本批报告是否全部修复」判定整单复测完成，会误置「复测完成」；
    现按**工单全部关联漏洞**是否闭环判定（与漏洞状态变更路径同口径）。"""
    system_name = "复测部分覆盖系统QQ"
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth, json={"system_name": system_name},
    )
    assert resp.status_code == 200, resp.text
    plan_id = resp.json()["id"]
    # 录入漏洞仅限认领者：先认领该工单
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/claim", headers=auth)
    assert resp.status_code == 200, resp.text
    resp = await client.post(
        "/api/v1/vulns", headers=auth,
        json={"title": "复测部分覆盖漏洞QQ", "level": 20, "testing_plan_id": plan_id},
    )
    assert resp.status_code == 200, resp.text

    doc = _build_report_docx(
        system_name, "http://10.6.6.6/qq", "10.6.6.6",
        sections=[("复测部分覆盖已修复漏洞QQ", "高危", "命令执行漏洞", True)],
    )
    _records, result = await _import_report(
        client, auth, "20260728复测部分覆盖系统QQ渗透测试复测报告.docx", doc,
    )
    assert result["created"] == 1

    resp = await client.get("/api/v1/testing-plans", headers=auth, params={"search": system_name})
    plans = resp.json()["items"]
    assert len(plans) == 1
    assert plans[0]["status"] == 50, "工单仍有未闭环漏洞，不得因单份报告全修复而置复测完成"
    assert plans[0]["retest_done_time"] == ""

async def test_report_import_full_rounds_flow(client: AsyncClient, auth: dict):
    """初测 + 两轮复测报告依次导入：漏洞归并为同一条、状态按 未修复→复测未修复→已修复 流转，
    复测轮次=2、计划复测完成(60)、三份报告均自动创建。"""
    system_name = "综合办公系统RR"
    target_url = "http://10.7.7.7/oarr"
    vuln_title = "越权-劳动合同变更审批RR"

    # 1) 初测报告：发现漏洞，未修复
    _records, result = await _import_report(
        client, auth, "20250917中移系统集成有限公司综合办公系统RR渗透测试报告.docx",
        _build_report_docx(system_name, target_url, "10.7.7.7",
                           sections=[(vuln_title, "高危", "逻辑漏洞", False)]),
    )
    assert result["created"] == 1
    plan = (await client.get("/api/v1/testing-plans", headers=auth,
                             params={"search": system_name})).json()["items"][0]
    assert plan["status"] == 30  # 初测完成（等待业务系统提交复测）
    vuln = await _find_vuln(client, auth, vuln_title)
    assert vuln["status"] == 10 and vuln["is_retest"] is False  # 初测发现漏洞，未修复

    # 2) 第一轮复测（无后缀）：部分未修复 → 漏洞置「复测未修复」(50+is_retest)
    #    标题带「（部分未修复）」后缀 + 汇总表状态「部分未修复」，验证归一化后与初测漏洞去重合并
    _records, result = await _import_report(
        client, auth, "20251011中移系统集成有限公司综合办公系统RR渗透测试复测报告.docx",
        _build_report_docx(system_name, target_url, "10.7.7.7",
                           sections=[(vuln_title + "（部分未修复）", "高危", "逻辑漏洞", False)],
                           status_texts=["部分未修复"]),
    )
    assert result["created"] == 1
    plan = (await client.get("/api/v1/testing-plans", headers=auth,
                             params={"search": system_name})).json()["items"][0]
    assert plan["status"] == 50  # 复测中
    vuln = await _find_vuln(client, auth, vuln_title)
    assert vuln["status"] == 50 and vuln["is_retest"] is True

    # 3) 第二轮复测（-1 后缀）：已修复 → 漏洞闭环，计划复测完成
    _records, result = await _import_report(
        client, auth, "20251011中移系统集成有限公司综合办公系统RR渗透测试复测报告-1.docx",
        _build_report_docx(system_name, target_url, "10.7.7.7",
                           sections=[(vuln_title, "高危", "逻辑漏洞", True)]),
    )
    assert result["created"] == 1

    # 漏洞始终只有一条，贯穿三轮
    resp = await client.get("/api/v1/vulns", headers=auth, params={"search": vuln_title})
    assert len(resp.json()["items"]) == 1
    vuln = resp.json()["items"][0]
    assert vuln["status"] == 60 and vuln["is_retest"] is True

    # 计划：复测完成(60)，两轮复测轮次记录
    plan = (await client.get("/api/v1/testing-plans", headers=auth,
                             params={"search": system_name})).json()["items"][0]
    assert plan["status"] == 60
    assert plan["retest_round_count"] == 2

    # 三份报告均自动创建，且都关联到同一工单
    resp = await client.get("/api/v1/reports", headers=auth, params={"search": system_name})
    reports = [r for r in resp.json()["items"] if r["project_name"] == system_name]
    assert len(reports) == 3, reports
    assert {r["testing_plan_id"] for r in reports} == {plan["id"]}

    # 报告导入复测补写源报告：两轮都以导入的初测报告为源（否则报告维度三态恒为「未发起复测」）
    first_report = next(r for r in reports if "复测" not in r["title"])
    assert {r["src_report_id"] for r in plan["retest_rounds"]} == {first_report["id"]}
    assert first_report["retest_state"] == "done"

async def test_import_confirm_into_report(client: AsyncClient, auth: dict):
    """Word 导入确认时关联报告：自动追加章节、漏洞进入修复中。"""
    resp = await client.post(
        "/api/v1/reports", headers=auth, json={"title": "导入关联报告", "sections": []},
    )
    report = resp.json()
    report_id = report["id"]

    resp = await client.get("/api/v1/imports/template", headers=auth)
    resp = await client.post(
        "/api/v1/imports", headers=auth,
        files={"file": ("关联报告样例.docx", BytesIO(resp.content),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    batch_id = resp.json()["id"]
    detail = await _wait_batch(client, auth, batch_id)
    assert detail["batch"]["status"] == "parsed", detail
    rec = detail["records"][0]

    # 不存在的报告被拒绝
    resp = await client.post(
        f"/api/v1/imports/{batch_id}/confirm", headers=auth,
        json={"record_ids": [rec["id"]], "report_id": 999999},
    )
    assert resp.status_code == 400

    resp = await client.post(
        f"/api/v1/imports/{batch_id}/confirm", headers=auth,
        json={"record_ids": [rec["id"]], "report_id": report_id},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["created"] == 1

    # 报告新增关联章节且编辑锁 revision +1（导出版本号 version 不受编辑影响）
    saved = (await client.get(f"/api/v1/reports/{report_id}", headers=auth)).json()
    assert len(saved["sections"]) == 1
    assert saved["sections"][0]["vul_id"] is not None
    assert saved["revision"] == report["revision"] + 1

    # 入库漏洞自动进入修复中
    vul = (await client.get(f"/api/v1/vulns/{saved['sections'][0]['vul_id']}", headers=auth)).json()
    assert vul["status"] == 50

async def test_import_report_author_from_testers(client: AsyncClient, auth: dict):
    """显式关联渗透测试工单导入报告：自动创建的报告作者取自工单测试人员姓名。"""
    system_name = "作者同步系统TA"
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={"system_name": system_name, "test_type": "渗透测试"},
    )
    assert resp.status_code == 200, resp.text
    plan_id = resp.json()["id"]
    # admin 认领工单 → testers=[管理员]
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/claim", headers=auth)
    assert resp.status_code == 200, resp.text
    assert resp.json()["testers"][0]["realname"] == "管理员"

    doc = _build_report_docx(
        system_name, "http://10.6.6.6/authorTA", "10.6.6.6",
        sections=[("报告作者同步漏洞TA", "高危", "逻辑漏洞", False)],
    )
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    resp = await client.post(
        "/api/v1/imports", headers=auth,
        files={"file": ("20260728作者同步系统TA渗透测试报告.docx", buf,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, resp.text
    batch_id = resp.json()["id"]
    detail = await _wait_batch(client, auth, batch_id)
    assert detail["batch"]["doc_kind"] == "report", detail
    rec_ids = [r["id"] for r in detail["records"]]

    resp = await client.post(
        f"/api/v1/imports/{batch_id}/confirm", headers=auth,
        json={"record_ids": rec_ids, "testing_plan_id": plan_id},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["created"] >= 1

    # 自动创建的报告作者 = 工单测试人员姓名（「、」拼接）
    resp = await client.get("/api/v1/reports", headers=auth, params={"search": system_name})
    reports = [r for r in resp.json()["items"] if r["project_name"] == system_name]
    assert reports, "未自动创建报告"
    report = (await client.get(f"/api/v1/reports/{reports[0]['id']}", headers=auth)).json()
    assert report["author"] == "管理员"
    assert report["testing_plan_id"] == plan_id

async def test_import_report_fields_and_auto_export(client: AsyncClient, auth: dict):
    """导入报告含参测人员与测试周期：映射系统账号关联工单、回填报告字段
    （作者/测试周期/测试账号/实际人天）、被测URL更新资产、报告时间取标题日期、
    自动生成可下载的导出记录（时间=报告日期 14:00）。"""
    # 创建参测人员对应系统账号（对应生产环境 admin/xna/xtz 三人场景）
    for username, realname in (("xna", "许宁安"), ("xtz", "薛田泽")):
        resp = await client.post(
            "/api/v1/users", headers=auth,
            json={"username": username, "password": "Tester@123", "realname": realname,
                  "email": "", "phone": "", "is_active": True},
        )
        assert resp.status_code == 200, resp.text

    system_name = "字段回填系统FB"
    doc = _build_report_docx(
        system_name, "http://10.7.7.7/fieldfb", "10.7.7.7",
        sections=[("字段回填漏洞FB", "中危", "信息泄露", False)],
        testers=["管理员", "许宁安", "薛田泽"], test_start="2026-06-30", test_end="2026-07-01",
        test_account="admin/Admin@123",
    )
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    resp = await client.post(
        "/api/v1/imports", headers=auth,
        files={"file": ("20260701字段回填系统FB渗透测试报告.docx", buf,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, resp.text
    batch_id = resp.json()["id"]
    detail = await _wait_batch(client, auth, batch_id)
    assert detail["batch"]["doc_kind"] == "report", detail
    rec_ids = [r["id"] for r in detail["records"]]

    resp = await client.post(
        f"/api/v1/imports/{batch_id}/confirm", headers=auth,
        json={"record_ids": rec_ids},
    )
    assert resp.status_code == 200, resp.text

    # 自动创建的工单已按姓名关联参测人员账号（admin=管理员、xna=许宁安、xtz=薛田泽）
    resp = await client.get("/api/v1/reports", headers=auth, params={"search": system_name})
    reports = [r for r in resp.json()["items"] if r["project_name"] == system_name]
    assert reports, "未自动创建报告"
    report = (await client.get(f"/api/v1/reports/{reports[0]['id']}", headers=auth)).json()
    plan = (await client.get(f"/api/v1/testing-plans/{report['testing_plan_id']}", headers=auth)).json()
    assert {u["realname"] for u in plan["testers"]} == {"管理员", "许宁安", "薛田泽"}
    assert {u["username"] for u in plan["testers"]} >= {"admin", "xna", "xtz"}

    # 报告字段回填：作者 / 测试周期 / 测试账号 / 实际人天 / 报告时间（取标题日期 2026-07-01 14:00）
    assert report["author"] == "管理员、许宁安、薛田泽"
    assert report["test_start"] == "2026-06-30"
    assert report["test_end"] == "2026-07-01"
    assert report["test_account"] == "admin/Admin@123"
    assert report["actual_mandays"] == 2  # 2026-06-30 ~ 2026-07-01
    assert report["create_time"].startswith("2026-07-01T14:00")
    # 自动导出成功：导出版本 +1；报告保持草稿（定稿仍由人工导出 Word 驱动）
    assert report["version"] == 2
    assert report["status"] == "draft"

    # 工单实际人天同步刷新（仅纳入初测报告）
    assert plan["actual_mandays"] == 2

    # 漏洞提交时间 = 报告时间（标题日期 14:00），保证按月统计口径一致
    resp = await client.get("/api/v1/vulns", headers=auth, params={"search": "字段回填漏洞FB"})
    vul = resp.json()["items"][0]
    assert vul["submit_time"].startswith("2026-07-01T14:00"), vul["submit_time"]

    # 被测系统 URL 自动更新到资产（internal_urls 去重）
    resp = await client.get("/api/v1/assets", headers=auth, params={"search": system_name})
    asset = [a for a in resp.json()["items"] if a["name"] == system_name][0]
    assert "http://10.7.7.7/fieldfb" in asset["internal_urls"]

    # 自动生成一条导出记录：时间=报告日期 14:00，含实际文件可下载
    resp = await client.get(f"/api/v1/reports/{report['id']}/exports", headers=auth)
    assert resp.status_code == 200, resp.text
    assert len(resp.json()) == 1, resp.text
    job = resp.json()[0]
    assert job["status"] == "done"
    assert job["create_time"].startswith("2026-07-01T14:00")
    assert job["finish_time"].startswith("2026-07-01T14:00")
    assert job["has_file"] is True
    dl = await client.get(f"/api/v1/reports/exports/{job['id']}/download", headers=auth)
    assert dl.status_code == 200, dl.text
    assert len(dl.content) > 0

async def test_import_report_public_urls_and_doc_time(client: AsyncClient, auth: dict):
    """多条被测系统 URL 分别录入资产（公网→public_urls、内网→internal_urls）；
    下载的 docx 封面与版本变更记录时间与导入报告时间一致。"""
    system_name = "多URL分类系统MU"
    doc = _build_report_docx(
        system_name, "https://www.a-mu.com\nhttps://www.b-mu.com\n10.30.30.30", "10.30.30.30",
        sections=[("多URL漏洞MU", "中危", "信息泄露", False)],
        testers=["管理员"], test_start="2026-07-10", test_end="2026-07-15",
    )
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    resp = await client.post(
        "/api/v1/imports", headers=auth,
        files={"file": ("20260715多URL分类系统MU渗透测试报告.docx", buf,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, resp.text
    batch_id = resp.json()["id"]
    detail = await _wait_batch(client, auth, batch_id)
    rec_ids = [r["id"] for r in detail["records"]]
    resp = await client.post(
        f"/api/v1/imports/{batch_id}/confirm", headers=auth,
        json={"record_ids": rec_ids},
    )
    assert resp.status_code == 200, resp.text

    # 公网 URL 分别录入 public_urls（tag=10 互联网），内网 IP 录入 internal_urls
    resp = await client.get("/api/v1/assets", headers=auth, params={"search": system_name})
    asset = [a for a in resp.json()["items"] if a["name"] == system_name][0]
    pub = {u["url"]: u["tag"] for u in (asset["public_urls"] or [])}
    assert pub.get("https://www.a-mu.com") == 10
    assert pub.get("https://www.b-mu.com") == 10
    assert "10.30.30.30" in (asset["internal_urls"] or [])

    # 下载自动导出的 docx：封面日期与版本变更记录 V1.0 日期均为导入报告时间 2026-07-15
    resp = await client.get("/api/v1/reports", headers=auth, params={"search": system_name})
    report = [r for r in resp.json()["items"] if r["project_name"] == system_name][0]
    jobs = (await client.get(f"/api/v1/reports/{report['id']}/exports", headers=auth)).json()
    assert len(jobs) == 1, jobs
    dl = await client.get(f"/api/v1/reports/exports/{jobs[0]['id']}/download", headers=auth)
    assert dl.status_code == 200, dl.text

    from docx import Document

    docx = Document(BytesIO(dl.content))
    texts = [p.text.strip() for p in docx.paragraphs if p.text.strip()]
    assert any("2026年07月15日" in t for t in texts), texts
    cells = [c.text.strip() for c in docx.tables[1].rows[2].cells]
    assert cells[0] == "2026-07-15", cells

async def test_import_report_section_and_list_order(client: AsyncClient, auth: dict):
    """导入报告章节顺序：确认入库按解析序号处理，报告章节 order、章节导航渲染顺序
    与 /vulns 默认视图顺序均与原报告序号一致（回归：PG 下无序处理导致章节乱序）。"""
    system_name = "章节顺序系统SX"
    titles = [f"章节顺序漏洞SX{i:02d}" for i in range(1, 6)]
    doc = _build_report_docx(
        system_name, "http://10.40.40.40/sx", "10.40.40.40",
        sections=[(t, "高危", "逻辑漏洞", False) for t in titles],
    )
    filename = "20260720章节顺序系统SX渗透测试报告.docx"
    records, _ = await _import_report(client, auth, filename, doc)
    # 解析记录本身按文档序号返回
    assert [r["title"] for r in records] == titles

    # 报告章节 order 与解析 seq 严格一致（章节导航按 order 渲染）
    resp = await client.get("/api/v1/reports", headers=auth, params={"search": system_name})
    report = [r for r in resp.json()["items"] if r["project_name"] == system_name][0]
    detail = (await client.get(f"/api/v1/reports/{report['id']}", headers=auth)).json()
    sections = sorted(detail["sections"], key=lambda s: s["order"])
    assert [s["title"] for s in sections] == titles
    assert [s["order"] for s in sections] == list(range(len(titles)))

    # /vulns 默认视图（不传 sort）：同 submit_time 批次按 id 升序 = 原报告序号
    plan_id = report["testing_plan_id"]
    resp = await client.get(
        "/api/v1/vulns", headers=auth, params={"testing_plan_id": plan_id, "size": 50},
    )
    items = resp.json()["items"]
    assert len(items) == len(titles)
    assert [v["title"] for v in items] == titles

    # 导入序号映射 {vul_id: seq}：漏洞 id 顺序与 seq 顺序一致，供流程抽屉同等级内按原报告序号排序
    resp = await client.get(f"/api/v1/testing-plans/{plan_id}/vuln-order", headers=auth)
    assert resp.status_code == 200, resp.text
    order_map = resp.json()
    assert len(order_map) == len(titles)
    assert [order_map[str(v["id"])] for v in items] == list(range(1, len(titles) + 1))

async def test_word_import_knowledge_backfill(client: AsyncClient, auth: dict):
    """Word 导入确认入库：描述/修复建议为空时自动套用知识库模板。"""
    # 准备 vul_type=15 的模板
    resp = await client.post(
        "/api/v1/knowledge", headers=auth,
        json={"vulnerability_name": "存储型XSS", "vul_type": 15, "severity_level": 20,
              "description_html": "<p>回填标准描述</p>",
              "harm_html": "<p>回填危害说明</p>", "solution_html": "<p>回填修复建议</p>"},
    )
    assert resp.status_code == 200, resp.text

    # 上传官方模板样例并等待解析
    resp = await client.get("/api/v1/imports/template", headers=auth)
    resp = await client.post(
        "/api/v1/imports", headers=auth,
        files={"file": ("知识库回填样例.docx", BytesIO(resp.content),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    batch_id = resp.json()["id"]
    detail = await _wait_batch(client, auth, batch_id)
    rec = detail["records"][0]

    # 清空描述/修复建议并改为模板类型，确认入库触发回填
    resp = await client.put(
        f"/api/v1/imports/records/{rec['id']}", headers=auth,
        json={"title": "知识库回填漏洞", "vul_type": 15,
              "description_html": "", "solution_html": ""},
    )
    assert resp.status_code == 200, resp.text
    resp = await client.post(
        f"/api/v1/imports/{batch_id}/confirm", headers=auth,
        json={"record_ids": [rec["id"]]},
    )
    assert resp.status_code == 200, resp.text

    resp = await client.get("/api/v1/vulns", headers=auth, params={"search": "知识库回填漏洞"})
    vul = resp.json()["items"][0]
    assert "回填标准描述" in vul["description_html"]
    assert "危害说明" in vul["description_html"]
    assert "回填危害说明" in vul["description_html"]
    assert vul["solution_html"] == "<p>回填修复建议</p>"

async def test_batch_import_confirm(client: AsyncClient, auth: dict):
    """批量确认入库：多批次统一关联工单 → 逐批确认并返回 report_ids；
    已确认批次重复关联跳过；单批失败隔离不影响其余批次；权限与参数校验。"""
    # 创建统一工单
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={"system_name": "批量确认系统", "test_type": "渗透测试"},
    )
    assert resp.status_code == 200, resp.text
    plan_id = resp.json()["id"]

    # 上传两份报告格式批次
    b1 = await _upload_report_batch(client, auth, "批量确认系统A", "20260801批量确认A渗透测试报告.docx")
    b2 = await _upload_report_batch(client, auth, "批量确认系统B", "20260801批量确认B渗透测试报告.docx")

    # 正常批量确认：统一关联工单，两个批次都成功，报告自动生成并挂到该工单
    resp = await client.post(
        "/api/v1/imports/batch-confirm", headers=auth,
        json={"batch_ids": [b1, b2], "testing_plan_id": plan_id},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["confirmed"] == 2
    assert body["skipped"] == 0
    assert body["failed"] == 0
    assert len(body["report_ids"]) == 2
    assert all(d["status"] == "confirmed" for d in body["details"])
    for rid in body["report_ids"]:
        report = (await client.get(f"/api/v1/reports/{rid}", headers=auth)).json()
        assert report["testing_plan_id"] == plan_id

    # 重复关联：已确认批次再批量确认（含 batch_ids 内重复）→ 全部 skipped
    resp = await client.post(
        "/api/v1/imports/batch-confirm", headers=auth,
        json={"batch_ids": [b1, b1, b2, b1]},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["confirmed"] == 0
    assert body["skipped"] == 2
    assert body["failed"] == 0
    assert all(d["status"] == "skipped" for d in body["details"])

    # 部分失败隔离：一个有效批次 + 一个不存在批次 → confirmed=1, failed=1，有效批次照常入库
    b3 = await _upload_report_batch(client, auth, "批量确认系统C", "20260801批量确认C渗透测试报告.docx")
    resp = await client.post(
        "/api/v1/imports/batch-confirm", headers=auth,
        json={"batch_ids": [b3, 999999]},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["confirmed"] == 1
    assert body["failed"] == 1
    assert len(body["report_ids"]) == 1
    failed_detail = next(d for d in body["details"] if d["status"] == "failed")
    assert failed_detail["batch_id"] == 999999
    assert failed_detail["detail"] == "导入批次不存在"

    # 空 batch_ids 拒绝
    resp = await client.post(
        "/api/v1/imports/batch-confirm", headers=auth, json={"batch_ids": []},
    )
    assert resp.status_code == 400

    # 非法工单拒绝
    resp = await client.post(
        "/api/v1/imports/batch-confirm", headers=auth,
        json={"batch_ids": [b1], "testing_plan_id": 999999},
    )
    assert resp.status_code == 400

    # 无 import:manage 权限用户 403
    resp = await client.post(
        "/api/v1/roles", headers=auth,
        json={"name": "无导入权限", "permissions": ["vuln:submit"], "remark": ""},
    )
    assert resp.status_code == 200, resp.text
    role_id = resp.json()["id"]
    resp = await client.post(
        "/api/v1/users", headers=auth,
        json={"username": "no_import_perm", "password": "Tester@123", "realname": "无导入",
              "email": "", "phone": "", "is_active": True, "role_id": role_id},
    )
    assert resp.status_code == 200, resp.text
    resp = await client.post(
        "/api/v1/auth/login", data={"username": "no_import_perm", "password": "Tester@123"},
    )
    auth2 = {"Authorization": f"Bearer {resp.json()['access_token']}"}
    resp = await client.post(
        "/api/v1/imports/batch-confirm", headers=auth2,
        json={"batch_ids": [b1]},
    )
    assert resp.status_code == 403

async def test_batch_confirm_chrono_order(client: AsyncClient, auth: dict):
    """批量确认时序：同一工单同一漏洞跨多份复测报告去重合并时，
    必须按报告日期从旧到新处理，保证最新报告的「已修复」最终生效而非被旧报告覆盖。"""
    system_name = "批量时序系统"
    target_url = "http://10.9.9.9/chrono"
    docx_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    async def upload(filename: str, fixed: bool) -> int:
        doc = _build_report_docx(
            system_name, target_url, "10.9.9.9",
            sections=[("XSS时序漏洞", "中危", "XSS跨站", fixed)],
        )
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        resp = await client.post(
            "/api/v1/imports", headers=auth,
            files={"file": (filename, buf, docx_mime)},
        )
        assert resp.status_code == 200, resp.text
        bid = resp.json()["id"]
        detail = await _wait_batch(client, auth, bid)
        assert detail["batch"]["status"] == "parsed", detail
        assert detail["batch"]["doc_kind"] == "report", detail
        return bid

    # 旧复测报告 XSS 未修复，新复测报告 XSS 已修复
    older = await upload("20260723批量时序系统渗透测试复测报告.docx", False)
    newer = await upload("20260731批量时序系统渗透测试复测报告.docx", True)

    # 以「新→旧」倒序批量确认（复现列表 newest-first 选中提交的顺序）
    resp = await client.post(
        "/api/v1/imports/batch-confirm", headers=auth,
        json={"batch_ids": [newer, older]},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["confirmed"] == 2

    # 同一工单同一标题应去重合并为一条漏洞，最终状态为「已修复」(60)
    vul = await _find_vuln(client, auth, "XSS时序漏洞")
    assert vul["status"] == 60, f"期望已修复(60)，实际 {vul['status']}"

async def test_import_level_mismatch_notice(client: AsyncClient, auth: dict):
    """风险问题汇总与风险问题详情等级不一致：入库取详情等级并暴露提醒明细。

    回归：旧实现只认汇总表，详情同行写法（「漏洞等级：高危」）被整体丢弃，
    匹配失败时静默回落中危，导致漏洞等级误判。
    """
    from docx import Document

    doc = Document()
    doc.add_heading("风险问题汇总", level=1)
    table = doc.add_table(rows=2, cols=4)
    for i, h in enumerate(("问题等级", "风险类型", "风险问题", "修复状态")):
        table.rows[0].cells[i].text = h
    table.rows[1].cells[0].text = "中危"
    table.rows[1].cells[1].text = "越权"
    table.rows[1].cells[2].text = "平行越权访问项目信息"
    table.rows[1].cells[3].text = "未修复"
    doc.add_heading("风险问题详情", level=1)
    doc.add_heading("平行越权访问项目信息（未修复）", level=3)
    doc.add_paragraph("测试状态：初测")
    doc.add_paragraph("漏洞等级：高危")
    doc.add_paragraph("漏洞描述：")
    doc.add_paragraph("可越权访问他人项目信息。")

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    resp = await client.post(
        "/api/v1/imports", headers=auth,
        files={"file": ("20260910等级不一致系统渗透测试报告.docx", buf,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, resp.text
    batch_id = resp.json()["id"]
    detail = await _wait_batch(client, auth, batch_id)
    rec = detail["records"][0]
    assert rec["level"] == 20  # 以风险问题详情的高危为准
    assert rec["level_source"] == "detail"
    assert rec["level_mismatch"] is True
    assert rec["level_summary_text"] == "中危"
    assert rec["level_detail_text"] == "高危"

    resp = await client.get(
        "/api/v1/imports/level-mismatches", headers=auth,
        params={"batch_ids": str(batch_id)},
    )
    assert resp.status_code == 200, resp.text
    items = resp.json()
    assert len(items) == 1, items
    assert items[0]["batch_id"] == batch_id
    assert items[0]["title"] == "平行越权访问项目信息"
    assert items[0]["level"] == 20
    assert items[0]["level_summary"] == 30
    assert items[0]["level_detail_text"] == "高危"

    # 不指定批次时同样可查到待入库的不一致记录（批量确认场景）
    resp = await client.get("/api/v1/imports/level-mismatches", headers=auth)
    assert any(x["batch_id"] == batch_id for x in resp.json())
