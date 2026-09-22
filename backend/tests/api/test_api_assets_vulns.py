"""API 集成测试：资产台账与漏洞库。

资产台账与漏洞库 CRUD：资产-漏洞生命周期、批量建漏洞、Excel
导入导出、列表排序、删除权限与引用清理、系统名检索、全局搜索。

原 tests/test_api.py（4626 行单体）按领域拆分而来：用例源码逐字未改，
仅按领域重新归位；共享夹具见 tests/conftest.py，共享 helper 见 tests/api/_helpers.py。
"""
from io import BytesIO
import pytest
from httpx import AsyncClient

from _helpers import _user_with_perms

pytestmark = pytest.mark.asyncio(loop_scope="session")


async def test_asset_and_vuln_lifecycle(client: AsyncClient, auth: dict):
    # 建资产（系统级：负责人 / URL 标签 / 端口服务对 / 中间件与数据库多条目）
    resp = await client.post(
        "/api/v1/assets", headers=auth,
        json={
            "name": "测试商城",
            "sub_system": "订单中心",
            "department": "电商事业部",
            "public_urls": [{"url": "https://shop.example.com", "tag": 10}],
            "internal_urls": ["http://10.0.0.8:8080"],
            "port_services": [
                {"port": "80", "service": "Web服务"},
                {"port": "443", "service": "HTTPS"},
            ],
            "middlewares": [{"name": "Nginx", "version": "1.24"}],
            "databases": [{"name": "MySQL", "version": "8.0"}],
            "owners": [{"name": "张三", "phone": "13800000000", "email": "zhangsan@example.com"}],
        },
    )
    assert resp.status_code == 200, resp.text
    asset = resp.json()
    asset_id = asset["id"]
    assert asset["owners"][0]["name"] == "张三"
    assert asset["public_urls"][0]["tag"] == 10
    assert asset["port_services"][1] == {"port": "443", "service": "HTTPS"}
    assert asset["middlewares"][0]["version"] == "1.24"
    assert asset["databases"][0] == {"name": "MySQL", "version": "8.0"}

    # 系统命名必填
    resp = await client.post("/api/v1/assets", headers=auth, json={"name": ""})
    assert resp.status_code == 422

    # 建漏洞（未修复 10），多对多关联资产
    resp = await client.post(
        "/api/v1/vulns", headers=auth,
        json={
            "title": "订单接口越权", "level": 20, "vul_type": 30,
            "affected_url": "https://shop.example.com/api/order",
            "description_html": "<p>横向越权读取他人订单</p>",
            "asset_ids": [asset_id],
        },
    )
    assert resp.status_code == 200, resp.text
    vul = resp.json()
    assert vul["status"] == 10
    assert vul["asset_ids"] == [asset_id]
    assert vul["assets"][0]["name"] == "测试商城"
    vul_id = vul["id"]

    # 按资产筛选漏洞
    resp = await client.get("/api/v1/vulns", headers=auth, params={"asset_id": asset_id})
    assert vul_id in [v["id"] for v in resp.json()["items"]]

    # 已关联漏洞的资产不能删除
    resp = await client.delete(f"/api/v1/assets/{asset_id}", headers=auth)
    assert resp.status_code == 400

    # 未修复可流转到：已忽略/暂不处理/修复中/复测中（复测中为修复状态冗余兜底路径）
    resp = await client.get(f"/api/v1/vulns/{vul_id}/transitions", headers=auth)
    assert {t["status"] for t in resp.json()} == {20, 35, 50, 55}

    # 10 -> 50 进入修复中，打 notice_time
    resp = await client.post(
        f"/api/v1/vulns/{vul_id}/transition", headers=auth,
        json={"status": 50, "comment": "已出报告"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["status"] == 50
    assert resp.json()["notice_time"] is not None

    # 非法流转 50 -> 60 应被拒绝（必须经过复测）
    resp = await client.post(
        f"/api/v1/vulns/{vul_id}/transition", headers=auth,
        json={"status": 60},
    )
    assert resp.status_code == 400

    # 50 -> 55 复测中 -> 60 已修复（终态，须携带复测内容）
    resp = await client.post(
        f"/api/v1/vulns/{vul_id}/transition", headers=auth, json={"status": 55},
    )
    assert resp.status_code == 200
    assert resp.json()["is_retest"] is True
    resp = await client.post(
        f"/api/v1/vulns/{vul_id}/transition", headers=auth,
        json={"status": 60, "comment": "复测通过", "retest_html": "<p>复测通过</p>"},
    )
    assert resp.status_code == 200
    assert resp.json()["fix_time"] is not None
    # 已修复为可回退状态：可重新打开为未修复(10)或直接重新复测(55)
    resp = await client.get(f"/api/v1/vulns/{vul_id}/transitions", headers=auth)
    assert {t["status"] for t in resp.json()} == {10, 55}

    # 日志包含创建与状态流转；展示名优先用户姓名（admin 姓名为「管理员」）
    resp = await client.get(f"/api/v1/vulns/{vul_id}/logs", headers=auth)
    logs = resp.json()
    actions = [log["action"] for log in logs]
    assert "创建漏洞" in actions
    assert len(actions) >= 2
    assert all("realname" in log for log in logs)
    assert any(log["realname"] == "管理员" for log in logs)

async def test_vuln_batch_create(client: AsyncClient, auth: dict):
    resp = await client.post(
        "/api/v1/assets", headers=auth, json={"name": "批量目标系统"},
    )
    asset_id = resp.json()["id"]

    resp = await client.post(
        "/api/v1/vulns/batch", headers=auth,
        json={
            "asset_ids": [asset_id],
            "vulns": [
                {"title": "批量漏洞A", "level": 20, "vul_type": 10},
                {"title": "批量漏洞B", "level": 30, "vul_type": 15},
                {"title": "批量漏洞C", "level": 40, "vul_type": 55},
            ],
        },
    )
    assert resp.status_code == 200, resp.text
    vulns = resp.json()
    assert len(vulns) == 3
    assert all(v["asset_ids"] == [asset_id] for v in vulns)

    # 无效资产ID被拒绝
    resp = await client.post(
        "/api/v1/vulns/batch", headers=auth,
        json={"asset_ids": [999999], "vulns": [{"title": "无效资产"}]},
    )
    assert resp.status_code == 400

async def test_asset_excel_import_export(client: AsyncClient, auth: dict):
    from openpyxl import Workbook, load_workbook

    # 构造导入文件：1 行合法 + 1 行缺系统命名
    wb = Workbook()
    ws = wb.active
    # 表头与模板一致（assets.py EXCEL_HEADERS，含「系统类型」列）
    ws.append(["系统命名*", "子系统名称", "部门", "系统类型", "公网URL", "内网URL",
               "开放端口与服务", "中间件", "数据库", "系统负责人", "状态", "备注"])
    ws.append(["Excel导入系统", "支付子系统", "金融部", "自有系统",
               "https://pay.example.com|互联网;https://oa.example.com|办公网",
               "http://192.168.1.10", "443:支付服务;8443:管理后台", "Tomcat/9.0", "Oracle/19c",
               "李四/13900000000/lisi@example.com;王五//wangwu@example.com",
               "线上", "备注内容"])
    ws.append(["", "无名子系统", "", "", "", "", "", "", "", "", "", ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = await client.post(
        "/api/v1/assets/import", headers=auth,
        files={"file": ("assets.xlsx", buf,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    result = resp.json()
    assert result["total"] == 2
    assert result["success"] == 1
    assert result["failed"] == 1
    assert "系统命名为必填项" in result["errors"][0]

    # 非 xlsx 拒绝
    resp = await client.post(
        "/api/v1/assets/import", headers=auth,
        files={"file": ("bad.txt", b"hello", "text/plain")},
    )
    assert resp.status_code == 400

    # 导入结果校验（多值字段解析）
    resp = await client.get("/api/v1/assets", headers=auth, params={"search": "Excel导入系统"})
    items = resp.json()["items"]
    assert len(items) == 1
    imported = items[0]
    assert imported["public_urls"] == [
        {"url": "https://pay.example.com", "tag": 10},
        {"url": "https://oa.example.com", "tag": 20},
    ]
    assert imported["port_services"] == [
        {"port": "443", "service": "支付服务"},
        {"port": "8443", "service": "管理后台"},
    ]
    assert imported["middlewares"] == [{"name": "Tomcat", "version": "9.0"}]
    assert imported["databases"] == [{"name": "Oracle", "version": "19c"}]
    assert len(imported["owners"]) == 2
    assert imported["owners"][1] == {"name": "王五", "phone": "", "email": "wangwu@example.com"}

    # 导出并回读
    resp = await client.get("/api/v1/assets/export", headers=auth, params={"search": "Excel导入系统"})
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.content))
    rows = list(wb.active.iter_rows(values_only=True))
    assert len(rows) == 2  # 表头 + 1 行数据
    assert rows[1][0] == "Excel导入系统"
    assert rows[1][3] == "自有系统"  # 系统类型列
    assert "互联网" in rows[1][4]  # 公网URL列
    assert rows[1][6] == "443:支付服务;8443:管理后台"
    assert rows[1][7] == "Tomcat/9.0"

    # 模板下载
    resp = await client.get("/api/v1/assets/import/template", headers=auth)
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"

async def test_vuln_list_sorting(client: AsyncClient, auth: dict):
    """漏洞列表按 level 升序排序结果非降序；sort 非法字段回退默认排序不报错。"""
    # 造三条不同等级的漏洞
    for lv in (40, 10, 30):
        resp = await client.post(
            "/api/v1/vulns", headers=auth,
            json={"title": f"排序用例漏洞-{lv}", "level": lv, "vul_type": 30},
        )
        assert resp.status_code == 200, resp.text

    resp = await client.get(
        "/api/v1/vulns", headers=auth, params={"sort": "level", "order": "asc", "size": 100},
    )
    assert resp.status_code == 200, resp.text
    levels = [v["level"] for v in resp.json()["items"]]
    assert levels == sorted(levels), levels

    # 降序
    resp = await client.get(
        "/api/v1/vulns", headers=auth, params={"sort": "level", "order": "desc", "size": 100},
    )
    levels = [v["level"] for v in resp.json()["items"]]
    assert levels == sorted(levels, reverse=True), levels

    # 非法排序字段回退默认排序（不报错）
    resp = await client.get(
        "/api/v1/vulns", headers=auth, params={"sort": "drop table", "order": "asc"},
    )
    assert resp.status_code == 200, resp.text

async def test_vuln_delete_permissions(client: AsyncClient, auth: dict):
    """单删漏洞权限：漏洞管理员（含 `*`）可删任意漏洞，其余账号与编辑口径一致
    （未关联工单：提交人本人；关联工单：认领者），无关账号 403。

    历史缺口：单删原先仅放行 vuln:manage，工单认领者（无该权限）在流程抽屉无法删除
    自己录入的漏洞。现放宽到认领者/提交人，同时**保留管理员原有的全量删除能力**
    （编辑侧是「管理员未认领也不放行」的严格认领口径，删除侧不跟随，
    否则前端按 `vuln:manage` 显示的删除按钮会点击 403）。
    """
    claimer = await _user_with_perms(client, auth, "del_claimer", ["vuln:submit", "special:manage"])
    bystander = await _user_with_perms(client, auth, "del_bystander", ["vuln:submit"])
    manager = await _user_with_perms(client, auth, "del_manager", ["vuln:manage"])

    # 场景 A：未关联工单的漏洞 —— 提交人本人可删，无关账号 403，漏洞管理员（admin/*）可删
    resp = await client.post(
        "/api/v1/vulns", headers=bystander, json={"title": "单删-路人提交", "level": 30},
    )
    vul_own = resp.json()
    resp = await client.delete(f"/api/v1/vulns/{vul_own['id']}", headers=claimer)
    assert resp.status_code == 403, resp.text
    resp = await client.delete(f"/api/v1/vulns/{vul_own['id']}", headers=bystander)
    assert resp.status_code == 200, resp.text

    resp = await client.post(
        "/api/v1/vulns", headers=bystander, json={"title": "单删-管理员删", "level": 30},
    )
    vul_admin = resp.json()
    resp = await client.delete(f"/api/v1/vulns/{vul_admin['id']}", headers=auth)
    assert resp.status_code == 200, resp.text

    # 场景 B：关联工单的漏洞 —— 认领者（无 vuln:manage）可删，未认领账号 403
    resp = await client.post(
        "/api/v1/testing-plans", headers=auth,
        json={"system_name": "单删权限工单", "test_type": "渗透测试"},
    )
    plan_id = resp.json()["id"]
    resp = await client.post(f"/api/v1/testing-plans/{plan_id}/claim", headers=claimer)
    assert resp.status_code == 200, resp.text

    rows: list[dict] = []
    for title in ("单删-工单漏洞", "单删-超管未认领删", "单删-管理员未认领删"):
        resp = await client.post(
            "/api/v1/vulns", headers=claimer,
            json={"title": title, "level": 20, "testing_plan_id": plan_id},
        )
        assert resp.status_code == 200, resp.text
        rows.append(resp.json())

    resp = await client.delete(f"/api/v1/vulns/{rows[0]['id']}", headers=bystander)
    assert resp.status_code == 403, resp.text
    resp = await client.delete(f"/api/v1/vulns/{rows[0]['id']}", headers=claimer)
    assert resp.status_code == 200, resp.text
    # 幂等：漏洞不存在时同样返回删除成功
    resp = await client.delete(f"/api/v1/vulns/{rows[0]['id']}", headers=claimer)
    assert resp.status_code == 200, resp.text

    # 场景 B2：回归护栏 —— 未认领该工单的漏洞管理员（含 *）仍可删除其漏洞（原行为）
    for hdr, row in ((auth, rows[1]), (manager, rows[2])):
        resp = await client.delete(f"/api/v1/vulns/{row['id']}", headers=hdr)
        assert resp.status_code == 200, resp.text

    resp = await client.delete(f"/api/v1/testing-plans/{plan_id}", headers=auth)
    assert resp.status_code == 200, resp.text

async def test_vuln_delete_cleans_remote_testing(client: AsyncClient, auth: dict):
    """删除被远程检测关联的漏洞后关联字段置空（生产 PG 强制外键，不清理会 500）。"""
    resp = await client.post(
        "/api/v1/vulns", headers=auth, json={"title": "单删-远程检测漏洞", "level": 30},
    )
    assert resp.status_code == 200, resp.text
    vul_id = resp.json()["id"]
    resp = await client.post(
        "/api/v1/remote-testings", headers=auth,
        json={"system_name": "单删权限远程检测", "vuln_id": vul_id},
    )
    assert resp.status_code == 200, resp.text
    rt_id = resp.json()["id"]
    assert resp.json()["vuln_id"] == vul_id

    resp = await client.delete(f"/api/v1/vulns/{vul_id}", headers=auth)
    assert resp.status_code == 200, resp.text
    resp = await client.get(
        "/api/v1/remote-testings", headers=auth, params={"search": "单删权限远程检测"},
    )
    assert resp.status_code == 200, resp.text
    rows = [r for r in resp.json()["items"] if r["id"] == rt_id]
    assert rows and rows[0]["vuln_id"] is None, rows

    resp = await client.delete(f"/api/v1/remote-testings/{rt_id}", headers=auth)
    assert resp.status_code == 200, resp.text

async def test_vuln_search_by_system_name(client: AsyncClient, auth: dict):
    """关键词搜索支持系统名称：标题/URL 命中之外，还应命中关联资产（系统）的名称。"""
    resp = await client.post(
        "/api/v1/assets", headers=auth, json={"name": "检索定位专用系统XQ"},
    )
    assert resp.status_code == 200, resp.text
    asset_id = resp.json()["id"]
    resp = await client.post(
        "/api/v1/vulns", headers=auth,
        json={
            "title": "与系统名毫无关联的标题XQ", "level": 30, "vul_type": 30,
            "affected_url": "http://xq.example.com/api",
            "description_html": "<p>xq</p>",
            "asset_ids": [asset_id],
        },
    )
    assert resp.status_code == 200, resp.text
    vul_id = resp.json()["id"]

    # 系统名称命中
    resp = await client.get(
        "/api/v1/vulns", headers=auth, params={"search": "检索定位专用系统XQ"},
    )
    assert resp.status_code == 200, resp.text
    assert vul_id in [v["id"] for v in resp.json()["items"]]

    # 标题命中依旧可用（回归确认）
    resp = await client.get(
        "/api/v1/vulns", headers=auth, params={"search": "与系统名毫无关联的标题XQ"},
    )
    assert resp.status_code == 200, resp.text
    assert vul_id in [v["id"] for v in resp.json()["items"]]

async def test_global_search(client: AsyncClient, auth: dict):
    """⌘K 全局搜索：空关键字返回空分组；按标题/名称模糊命中漏洞与资产。"""
    resp = await client.get("/api/v1/search", headers=auth, params={"q": "  "})
    assert resp.status_code == 200
    assert resp.json() == {"vulns": [], "assets": [], "plans": [], "reports": []}

    # 造数据：资产 + 挂在其上的漏洞
    resp = await client.post("/api/v1/assets", headers=auth, json={"name": "搜索目标系统Alpha"})
    assert resp.status_code == 200, resp.text
    asset_id = resp.json()["id"]
    resp = await client.post(
        "/api/v1/vulns/batch", headers=auth,
        json={
            "asset_ids": [asset_id],
            "vulns": [{"title": "搜索专用SQL注入漏洞XYZ", "level": 20, "vul_type": 10}],
        },
    )
    assert resp.status_code == 200, resp.text

    resp = await client.get("/api/v1/search", headers=auth, params={"q": "XYZ"})
    assert resp.status_code == 200
    body = resp.json()
    assert [v["title"] for v in body["vulns"]] == ["搜索专用SQL注入漏洞XYZ"]
    assert body["assets"] == []

    # 按资产名搜索命中资产分区
    resp = await client.get("/api/v1/search", headers=auth, params={"q": "Alpha"})
    assert resp.status_code == 200
    body = resp.json()
    assert any(a["name"] == "搜索目标系统Alpha" for a in body["assets"])
    assert body["vulns"] == []
