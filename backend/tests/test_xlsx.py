"""上传 xlsx 解析统一封装（core/xlsx.load_xlsx）回归测试。

背景缺陷（2026-09-17 审计 C-3）：资产/工单导入口曾用 `except Exception` 把一切异常
统一转成 400「Excel 文件解析失败」，连服务端自身缺陷也会被伪装成「用户文件有问题」，
排障时丢失堆栈。现统一为：
- 可预期的「文件不可解析」异常（坏 ZIP / 非 xlsx / 内部 XML 损坏）→ HTTPException(400)；
- 其余异常继续上抛，由全局兜底处理并记录堆栈。
"""
from io import BytesIO

import pytest
from fastapi import HTTPException


def test_load_xlsx_rejects_non_zip_bytes():
    from app.core.xlsx import load_xlsx

    with pytest.raises(HTTPException) as exc:
        load_xlsx(b"definitely not an excel file")
    assert exc.value.status_code == 400
    assert "Excel 文件解析失败" in str(exc.value.detail)


def test_load_xlsx_parses_valid_workbook():
    from openpyxl import Workbook

    from app.core.xlsx import load_xlsx

    buf = BytesIO()
    Workbook().save(buf)
    wb = load_xlsx(buf.getvalue())
    assert wb is not None
    assert list(wb.sheetnames)


def test_load_xlsx_does_not_swallow_unexpected_errors(monkeypatch):
    """服务端缺陷（非用户文件问题）必须上抛，不得统一转成 400。"""
    import app.core.xlsx as mod

    def boom(*args, **kwargs):
        raise RuntimeError("服务端缺陷")

    monkeypatch.setattr(mod, "load_workbook", boom)
    with pytest.raises(RuntimeError):
        mod.load_xlsx(b"any bytes")
