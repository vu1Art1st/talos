"""DOCX 解析行为契约（批次 E-2 Phase 0）。

用途：把 `services/docx_parser.py` 当前**依赖 python-docx 语义**的行为固化成可执行契约，
供后续「lxml.iterparse 自写 reader」做差分对照。契约逐条对应 `docs/ROADMAP.md` 6.4.2。

原则：这里锁定的是**当前行为**，不是「更正确的行为」。已知的数据丢失（如 `w:hyperlink`
内的 run 不被 `p.runs` 收录）属于既有事实，改动它属于行为变更，必须另立专项——因此这里
显式断言其存在，防止被顺手「修复」。
"""
import base64
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from app.services.docx_parser import parse_any_docx, parse_docx

# 1x1 透明 PNG（与 test_parser.py 同一份）
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _field_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    for i, (label, value) in enumerate(rows):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
    return table


def _base_rows():
    return [
        ("漏洞名称", "契约用例-后台SQL注入"),
        ("漏洞等级", "高危"),
        ("漏洞类型", "SQL注入漏洞"),
        ("漏洞链接", "https://example.com/api/login"),
        ("漏洞描述", "描述正文"),
        ("漏洞证明", "证明正文"),
        ("修复建议", "修复建议正文"),
    ]


def _add_hyperlink(paragraph, text, rid="rIdContract"):
    """手工注入 `w:hyperlink`（python-docx 无公开 API）。"""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def test_contract_template_fields_and_determinism(tmp_path):
    """契约 1/2/4：模板字段抽取稳定，且同一文档重复解析结果一致。"""
    path = tmp_path / "contract_template.docx"
    doc = Document()
    _field_table(doc, _base_rows())
    doc.save(str(path))
    img = tmp_path / "img"

    first = parse_docx(str(path), str(img), "/x")
    second = parse_docx(str(path), str(img), "/x")

    assert first == second
    assert len(first) == 1
    rec = first[0]
    assert rec["title"] == "契约用例-后台SQL注入"
    assert rec["level"] == 20
    assert rec["affected_url"] == "https://example.com/api/login"


def test_contract_run_styles_bold_italic_and_explicit_false(tmp_path):
    """契约 5：只认直接 `w:rPr/w:b|w:i`；`w:val="0"` 视为不加粗。"""
    path = tmp_path / "contract_styles.docx"
    doc = Document()
    table = _field_table(doc, _base_rows())
    cell = table.rows[4].cells[1]  # 漏洞描述
    para = cell.add_paragraph()
    para.add_run("加粗段").bold = True
    para.add_run("斜体段").italic = True
    para.add_run("显式不加粗").bold = False
    doc.save(str(path))

    html = parse_docx(str(path), str(tmp_path / "img"), "/x")[0]["description_html"]

    assert "<strong>加粗段</strong>" in html
    assert "<em>斜体段</em>" in html
    assert "显式不加粗" in html
    assert "<strong>显式不加粗</strong>" not in html


def test_contract_image_extraction_bytes_and_order(tmp_path):
    """契约 6/7：run 内 `a:blip` 落盘为 uuid 文件名，字节与原文一致，且保持段内顺序。"""
    path = tmp_path / "contract_image.docx"
    doc = Document()
    table = _field_table(doc, _base_rows())
    cell = table.rows[5].cells[1]  # 漏洞证明
    para = cell.add_paragraph()
    para.add_run("前置文字")
    para.add_run().add_picture(BytesIO(_PNG), width=Inches(0.2))
    para.add_run("后置文字")
    doc.save(str(path))
    img_dir = tmp_path / "img"

    html = parse_docx(str(path), str(img_dir), "/storage/uploads/images")[0]["reproduce_html"]

    saved = list(img_dir.glob("*"))
    assert len(saved) == 1
    assert saved[0].read_bytes() == _PNG
    assert f'<img src="/storage/uploads/images/{saved[0].name}">' in html
    assert html.index("前置文字") < html.index("<img") < html.index("后置文字")


def test_contract_hyperlink_runs_are_excluded(tmp_path):
    """契约 3（已知数据丢失，刻意锁定）：`w:hyperlink` 内的 run 不在 `p.runs` 中。"""
    path = tmp_path / "contract_hyperlink.docx"
    doc = Document()
    table = _field_table(doc, _base_rows())
    cell = table.rows[4].cells[1]  # 漏洞描述
    para = cell.add_paragraph()
    para.add_run("可见文字")
    _add_hyperlink(para, "超链接文字")
    doc.save(str(path))

    html = parse_docx(str(path), str(tmp_path / "img"), "/x")[0]["description_html"]

    assert "可见文字" in html
    assert "超链接文字" not in html, "超链接内文字当前被丢弃；补回属行为变更，需另立专项"


def test_contract_nested_table_is_not_a_record(tmp_path):
    """契约 1：`doc.tables` 只含 body 直接子表，嵌套表不产生额外记录。"""
    path = tmp_path / "contract_nested.docx"
    doc = Document()
    table = _field_table(doc, _base_rows())
    nested = table.rows[5].cells[1].add_table(rows=1, cols=2)
    nested.rows[0].cells[0].text = "漏洞名称"
    nested.rows[0].cells[1].text = "嵌套表不应成为独立记录"
    doc.save(str(path))

    records = parse_docx(str(path), str(tmp_path / "img"), "/x")

    assert len(records) == 1
    assert records[0]["title"] == "契约用例-后台SQL注入"


def test_contract_report_detection_and_image(tmp_path):
    """契约 9/10：报告格式识别依赖段落样式名，且报告详情同样解析图片。"""
    path = tmp_path / "20260101契约系统渗透测试报告.docx"
    doc = Document()
    doc.add_heading("契约系统", level=0)
    doc.add_paragraph("渗透测试报告")
    summary = doc.add_table(rows=2, cols=4)
    for i, head in enumerate(("问题等级", "风险类型", "风险问题", "修复状态")):
        summary.rows[0].cells[i].text = head
    for i, value in enumerate(("高危", "SQL注入漏洞", "详情标题", "未修复")):
        summary.rows[1].cells[i].text = value
    doc.add_heading("风险问题详情", level=1)
    doc.add_heading("详情标题", level=3)
    doc.add_paragraph("漏洞等级：高危")
    doc.add_paragraph("漏洞证明")
    para = doc.add_paragraph()
    para.add_run().add_picture(BytesIO(_PNG), width=Inches(0.2))
    doc.save(str(path))
    img_dir = tmp_path / "img"

    kind, meta, records = parse_any_docx(
        str(path), str(img_dir), "/storage/uploads/images", path.name
    )

    assert kind == "report"
    assert meta["system_name"] == "契约系统"
    assert len(records) == 1
    assert records[0]["level"] == 20
    assert len(list(img_dir.glob("*"))) == 1
